import io
import json
import queue
import tempfile
import unittest
import wave
import zipfile
from array import array
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, call, patch

from datetime import datetime

from kunstlopassistenten import (
    App,
    add_missing_segments,
    auto_group_size,
    build_startliste,
    class_sort_key,
    compute_waveform,
    detect_audio_format,
    export_setup_package,
    extract_setup_package,
    find_music_for_row,
    fit_columns,
    generate_startliste_excel,
    generate_startliste_pdf,
    generate_vlc_playlist,
    get_audio_duration,
    group_overflows,
    layout_items,
    load_participants_from_spond,
    new_participant,
    ovelse_from_music,
    parse_duration_mmss,
    parse_music_path,
    parse_spond_date,
    place_water_pauses,
    split_segments,
    validate_setup,
    water_pause_positions,
)


class FakeVariable:
    def __init__(self, value):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class FakeTree:
    def __init__(self, columns, displaycolumns=("#all",)):
        self.columns = tuple(columns)
        self.displaycolumns = tuple(displaycolumns)
        self.config = {}
        self.selected = None
        self.focused = None

    def __getitem__(self, key):
        if key == "columns":
            return self.columns
        if key == "displaycolumns":
            return self.displaycolumns
        raise KeyError(key)

    def configure(self, **kwargs):
        self.config.update(kwargs)
        if "displaycolumns" in kwargs:
            self.displaycolumns = tuple(kwargs["displaycolumns"])

    def identify(self, region, _x, _y):
        return "cell" if region == "region" else ""

    def identify_column(self, _x):
        return "#1"

    def identify_row(self, _y):
        return "I001"

    def index(self, _item_id):
        return 0

    def selection_set(self, item_id):
        self.selected = item_id

    def focus(self, item_id):
        self.focused = item_id

    def get_children(self):
        return ("I001",)


class AudioFormatTests(unittest.TestCase):
    def test_detects_supported_audio_signatures(self):
        self.assertEqual(detect_audio_format(b"RIFF\x00\x00\x00\x00WAVE"), ("WAV", ".wav"))
        self.assertEqual(
            detect_audio_format(b"\x00\x00\x00\x1cftypM4A "),
            ("M4A", ".m4a"),
        )
        self.assertEqual(detect_audio_format(b"ID3\x04\x00"), ("MP3", ".mp3"))
        self.assertEqual(detect_audio_format(b"\xff\xfb\x90\x04"), ("MP3", ".mp3"))
        self.assertEqual(detect_audio_format(b"ikke lyd"), ("Ukjent", ""))

    def test_reads_wav_duration(self):
        data = io.BytesIO()
        with wave.open(data, "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(8000)
            audio.writeframes(b"\x00\x00" * 8000)

        duration = get_audio_duration(data.getvalue(), "WAV")

        self.assertAlmostEqual(duration, 1.0, places=3)

    def test_all_supplied_music_files_have_duration(self):
        music_zip = (
            Path(__file__).resolve().parent
            / "Musikk_02-09-2026_KUNSTLØP Oppvisning Bergen.zip"
        )
        if not music_zip.exists():
            self.skipTest("Musikk-ZIP er ikke tilgjengelig.")

        missing = []
        with zipfile.ZipFile(music_zip, "r") as archive:
            for entry in archive.infolist():
                if entry.is_dir():
                    continue
                data = archive.read(entry)
                audio_format, _suffix = detect_audio_format(data)
                if not get_audio_duration(data, audio_format):
                    missing.append(entry.filename)

        self.assertEqual(missing, [])


class ScheduleTests(unittest.TestCase):
    def runners(self, count):
        return [new_participant(f"Løper{i}", "Test", "Loddefjord IL") for i in range(1, count + 1)]

    def test_groups_warmup_and_times(self):
        items = layout_items(self.runners(3), 2, 2, 0)
        entries = build_startliste(
            items, 190, datetime(2000, 1, 1, 8, 15), ready_seconds=30, warmup_seconds=240
        )

        self.assertEqual(
            [(e["navn"], e["start"], e["end"], e["duration"]) for e in entries],
            [
                ("Oppvarmingsgruppe 1", "08:15:00", "08:19:00", "00:04:00"),
                ("Løper1 Test", "08:19:00", "08:22:40", "00:03:40"),
                ("Løper2 Test", "08:22:40", "08:26:20", "00:03:40"),
                ("Oppvarmingsgruppe 2", "08:26:20", "08:30:20", "00:04:00"),
                ("Løper3 Test", "08:30:20", "08:34:00", "00:03:40"),
            ],
        )
        self.assertEqual([e["nr"] for e in entries if not e["is_group"]], [1, 2, 3])

    def test_cancelled_runners_are_hidden_last(self):
        rows = self.runners(3)
        rows[0]["Påmelding"] = "Avmeldt"

        items = layout_items(rows, 1, 8, 0)

        self.assertIs(items[-1], rows[0])
        entries = build_startliste(items, 60, datetime(2000, 1, 1, 8, 0))
        self.assertEqual([e["navn"] for e in entries if e["kind"] == "run"], ["Løper2 Test", "Løper3 Test"])

    def test_water_pauses_between_groups(self):
        self.assertEqual(water_pause_positions(4, 1), [2])
        self.assertEqual(water_pause_positions(4, 2), [1, 3])
        self.assertEqual(water_pause_positions(1, 1), [])

        items = layout_items(self.runners(4), 2, 2, 1)
        entries = build_startliste(items, 60, datetime(2000, 1, 1, 8, 0), water_seconds=900)

        self.assertEqual([e["kind"] for e in entries],
                         ["warmup", "run", "run", "pause", "warmup", "run", "run"])
        self.assertEqual((entries[3]["start"], entries[3]["end"]), ("08:02:00", "08:17:00"))

    def test_no_warmup_groups(self):
        items = layout_items(self.runners(4), 0, 8, 1)

        self.assertEqual([i.get("Kind", "run") for i in items], ["run", "run", "pause", "run", "run"])

    def test_group_count_splits_runners(self):
        self.assertEqual(auto_group_size(7, 2), 4)
        self.assertEqual(auto_group_size(9, 3), 3)

        items = layout_items(self.runners(7), 3, 2, 0)
        self.assertEqual([i.get("Kind", "run") for i in items],
                         ["warmup", "run", "run", "warmup", "run", "run", "warmup", "run", "run", "run"])

        items = layout_items(self.runners(3), 1, 1, 0)
        self.assertEqual([i.get("Kind", "run") for i in items], ["warmup", "run", "run", "run"])

    def test_custom_run_length(self):
        rows = self.runners(2)
        rows[0]["LopLengde"] = 150
        entries = build_startliste(layout_items(rows, 0, 8, 0), 190, datetime(2000, 1, 1, 8, 0),
                                   ready_seconds=30)

        self.assertEqual([(e["run_length"], e["duration"], e["custom_length"]) for e in entries],
                         [("00:02:30", "00:03:00", True), ("00:03:10", "00:03:40", False)])

    def test_detects_group_with_too_many_runners(self):
        items = layout_items(self.runners(4), 2, 2, 0)
        self.assertEqual(group_overflows(items, 2), {})

        runner = items.pop(4)  # første utøver i gruppe 2 flyttes inn i gruppe 1
        items.insert(1, runner)

        self.assertEqual(group_overflows(items, 2), {id(items[0]): 3})

    def test_excel_has_original_sheet_and_sheet_with_duration(self):
        import openpyxl

        items = layout_items(self.runners(3), 2, 2, 0)
        entries = build_startliste(items, 190, datetime(2000, 1, 1, 8, 15), ready_seconds=30,
                                   warmup_seconds=240)
        with tempfile.TemporaryDirectory() as temp_dir:
            path = str(Path(temp_dir) / "startliste.xlsx")
            self.assertTrue(generate_startliste_excel(entries, path, "Tittel", Mock()))
            wb = openpyxl.load_workbook(path)

        self.assertEqual(wb.sheetnames, ["Startliste", "Med tidsforbruk"])
        plain, timed = wb["Startliste"], wb["Med tidsforbruk"]
        self.assertEqual([c.value for c in plain[3]],
                         ["Nr.", "Start", None, "Slutt", "Navn", "Klasse", "Program", "Klubb"])
        self.assertEqual([c.value for c in plain[4]][1:4], ["08:15:00", "-", "08:19:00"])
        self.assertEqual(plain["E4"].value, "Oppvarmingsgruppe 1")
        self.assertEqual(plain["B7"].value, "ca. 08:26:20")
        self.assertEqual([c.value for c in timed[3]],
                         ["Nr.", "Start", None, "Slutt", "Lengde på løp", "Sammenlagt tidsforbruk",
                          "Navn", "Klasse", "Program", "Klubb"])
        self.assertEqual([c.value for c in timed[5]], [1, "08:19:00", "-", "08:22:40", "00:03:10",
                                                       "00:03:40", "Løper1 Test", None, "Friløp",
                                                       "Loddefjord IL"])

    def test_moved_water_pause_keeps_its_place_until_count_changes(self):
        items = layout_items(self.runners(4), 2, 2, 1)
        pause = items.pop(3)
        items.insert(1, pause)

        entries = build_startliste(items, 60, datetime(2000, 1, 1, 8, 0), water_seconds=600)
        self.assertEqual(entries[1]["kind"], "pause")

        replaced = place_water_pauses(items, 1)
        self.assertEqual([i.get("Kind", "run") for i in replaced],
                         ["warmup", "run", "run", "pause", "warmup", "run", "run"])


class AudioCacheTests(unittest.TestCase):
    def make_app(self, temp_dir, data):
        music_zip = Path(temp_dir) / "Musikk.zip"
        with zipfile.ZipFile(music_zip, "w") as archive:
            archive.writestr("mappe/sang.mp3", data)

        app = App.__new__(App)
        app.music_zip = music_zip
        app.music_cache_dir = Path(temp_dir) / "cache"
        app.music_cache_dir.mkdir()
        app.music_file_info = {
            "mappe/sang.mp3": {"format": "MP3", "suffix": ".mp3"}
        }
        return app

    def test_reuses_existing_cache_file_without_rewriting(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            app = self.make_app(temp_dir, b"ID3-lyddata")
            first_path = app.get_cached_mp3_path("mappe/sang.mp3")

            with patch.object(Path, "write_bytes") as write_bytes:
                second_path = app.get_cached_mp3_path("mappe/sang.mp3")

            self.assertEqual(second_path, first_path)
            write_bytes.assert_not_called()

    def test_changed_audio_content_gets_new_cache_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            app = self.make_app(temp_dir, b"ID3-forste-versjon")
            first_path = app.get_cached_mp3_path("mappe/sang.mp3")
            with zipfile.ZipFile(app.music_zip, "w") as archive:
                archive.writestr("mappe/sang.mp3", b"ID3-andre-versjon")

            second_path = app.get_cached_mp3_path("mappe/sang.mp3")

            self.assertNotEqual(second_path, first_path)
            self.assertEqual(second_path.read_bytes(), b"ID3-andre-versjon")


class TableInteractionTests(unittest.TestCase):
    def test_resolves_column_after_columns_are_hidden(self):
        app = App.__new__(App)
        app.tree = FakeTree(
            ("startnummer", "musikknavn"),
            displaycolumns=("musikknavn",),
        )

        column_name = app.get_tree_column_name(10)

        self.assertEqual(column_name, "musikknavn")

    def test_updates_visible_columns(self):
        app = App.__new__(App)
        app.table_columns = ("startnummer", "musikknavn")
        app.table_column_vars = {
            "startnummer": FakeVariable(False),
            "musikknavn": FakeVariable(True),
        }
        app.tree = FakeTree(app.table_columns)

        app.update_visible_table_columns("startnummer")

        self.assertEqual(app.tree.displaycolumns, ("musikknavn",))

    def test_prevents_hiding_every_column(self):
        app = App.__new__(App)
        app.table_columns = ("startnummer",)
        app.table_column_vars = {"startnummer": FakeVariable(False)}
        app.tree = FakeTree(app.table_columns)

        with patch("kunstlopassistenten.messagebox.showinfo") as showinfo:
            app.update_visible_table_columns("startnummer")

        self.assertTrue(app.table_column_vars["startnummer"].get())
        showinfo.assert_called_once()

    def test_checkbox_requires_confirmation_and_marks_row_completed(self):
        app = App.__new__(App)
        app.tree = FakeTree(("lop_gjennomfort",))
        app.rows = [{"NavnFraIsonen": "Testløper"}]
        app.root = object()
        app.log = Mock()
        app.refresh_table = Mock()
        event = SimpleNamespace(x=10, y=10)

        with patch("kunstlopassistenten.messagebox.askyesno", return_value=True) as askyesno:
            result = app.on_tree_checkbox_click(event)

        self.assertEqual(result, "break")
        self.assertTrue(app.rows[0]["LopGjennomfort"])
        askyesno.assert_called_once()
        app.refresh_table.assert_called_once()

    def test_checkbox_removal_can_be_cancelled(self):
        app = App.__new__(App)
        app.tree = FakeTree(("lop_gjennomfort",))
        app.rows = [{"NavnFraIsonen": "Testløper", "LopGjennomfort": True}]
        app.root = object()
        app.log = Mock()
        app.refresh_table = Mock()
        event = SimpleNamespace(x=10, y=10)

        with patch("kunstlopassistenten.messagebox.askyesno", return_value=False) as askyesno:
            result = app.on_tree_checkbox_click(event)

        self.assertEqual(result, "break")
        self.assertTrue(app.rows[0]["LopGjennomfort"])
        askyesno.assert_called_once()
        app.refresh_table.assert_not_called()


class RegressionTests(unittest.TestCase):
    def runner(self, given, family="Test", **fields):
        row = new_participant(given, family)
        row.update(fields)
        return row

    def test_rejects_negative_and_invalid_durations(self):
        self.assertEqual(parse_duration_mmss("3:10"), 190)
        self.assertEqual(parse_duration_mmss("1:02:03"), 3723)
        for value in ("-1:00", "1:-05", "3:75", "1:60:00", "abc", "1"):
            self.assertIsNone(parse_duration_mmss(value), value)

    def test_stale_selection_after_list_is_replaced(self):
        app = App.__new__(App)
        app.schedule_tree = SimpleNamespace(selection=lambda: ("i5",), focus=lambda: "i5")
        app.items = [self.runner("A")]

        self.assertIsNone(app.selected_item())

    def test_competition_mode_dims_menus_and_other_tabs(self):
        app = App.__new__(App)
        app.log = Mock()
        app.menu_buttons = [Mock(), Mock()]
        app.music_tab = "stevne"
        app.notebook = Mock(tabs=Mock(return_value=("tidsplan", "stevne")))
        app.music_bar, app.music_hint, app.tree, app.mode_button = Mock(), Mock(), Mock(), Mock()

        app.set_competition_mode(True)

        app.menu_buttons[0].configure.assert_called_with(state="disabled")
        app.notebook.tab.assert_called_once_with("tidsplan", state="disabled")
        app.notebook.select.assert_called_once_with("stevne")
        app.music_bar.grid_remove.assert_called_once()
        app.mode_button.configure.assert_called_with(text="Redigeringsmodus")

        app.set_competition_mode(False)

        app.notebook.tab.assert_called_with("tidsplan", state="normal")
        app.music_bar.grid.assert_called_once()
        app.mode_button.configure.assert_called_with(text="Stevnemodus")

    def test_single_actions_use_focused_row_in_multi_selection(self):
        app = App.__new__(App)
        app.items = [self.runner("A"), self.runner("B"), self.runner("C")]
        app.schedule_tree = SimpleNamespace(selection=lambda: ("i0", "i2"), focus=lambda: "i2")

        self.assertIs(app.selected_item(), app.items[2])
        self.assertEqual([r["GivenName"] for r in app.selected_rows()], ["A", "C"])

    def test_short_program_added_later_creates_new_run(self):
        app = App.__new__(App)
        app.log = Mock()
        row = self.runner("Anna")
        app.items = split_segments([row])
        app.music_index = {"FS.mp3": parse_music_path("Anna_Test_FS.mp3")}
        app.apply_music_segments()
        self.assertEqual([(r["Segment"], r.get("Program")) for r in app.rows], [("FS", "")])

        app.music_index["SP.mp3"] = parse_music_path("Anna_Test_SP.mp3")
        app.apply_music_segments()

        self.assertEqual([(r["Segment"], r.get("Program")) for r in app.rows],
                         [("SP", "Kortprogram"), ("FS", "Friløp")])

    def test_imported_ovelse_is_not_changed_by_music_files(self):
        app = App.__new__(App)
        app.log = Mock()
        app.items = split_segments([self.runner("Anna", Ovelse="Friløp")])
        app.music_index = {"SP.mp3": parse_music_path("Anna_Test_SP.mp3")}

        app.apply_music_segments()

        self.assertEqual([r["Segment"] for r in app.rows], ["FS"])

    def test_ambiguous_music_match_is_marked_uncertain(self):
        index = {f: parse_music_path(f) for f in (
            "Stevne/Junior_Jenter/FS/Anna_Berg_FS.mp3",
            "Stevne/Senior_Damer/FS/Anna_Berg_FS.mp3",
        )}
        row = self.runner("Anna", "Berg", Segment="FS", Klasse="Ukjent klasse")

        matched, uncertain = find_music_for_row(row, index, set())

        self.assertTrue(matched)
        self.assertTrue(uncertain)

    def test_music_match_uses_class_and_is_certain(self):
        index = {f: parse_music_path(f) for f in (
            "Stevne/Junior_Jenter/FS/Anna_Berg_FS.mp3",
            "Stevne/Senior_Damer/FS/Anna_Berg_FS.mp3",
        )}
        row = self.runner("Anna", "Berg", Segment="FS", Klasse="Senior")

        self.assertEqual(find_music_for_row(row, index, set()),
                         ("Stevne/Senior_Damer/FS/Anna_Berg_FS.mp3", False))

    def test_exact_names_are_matched_before_fuzzy_names(self):
        app = App.__new__(App)
        app.music_files = ["S/Klasse/FS/Anna_Berg_Hansen_FS.mp3"]
        app.music_index = {f: parse_music_path(f) for f in app.music_files}
        app.music_file_info, app.music_durations, app.music_status = {}, {}, {}
        fuzzy_row = self.runner("Anna", "Berg", Segment="FS")
        exact_row = self.runner("Anna", "Berg Hansen", Segment="FS")

        app.match_rows([fuzzy_row, exact_row], set())

        self.assertEqual(exact_row["MusikkFil"], app.music_files[0])
        self.assertEqual(fuzzy_row["MusikkFil"], "")

    def test_late_short_program_info_adds_free_skate(self):
        row = self.runner("Anna")
        items = [{"Kind": "warmup"}] + split_segments([row])
        row["Ovelse"] = "Kortprogram+Friløp"

        self.assertEqual(add_missing_segments(items), 1)
        self.assertEqual([(i.get("Segment"), i.get("Program")) for i in items[1:]],
                         [("SP", "Kortprogram"), ("FS", "Friløp")])
        self.assertEqual(add_missing_segments(items), 0)

    def test_class_order_follows_skating_federation_list(self):
        classes = ["", "Ukjent", "Junior Kvinner", "Cubs u/13", "Cubs", "Oppvisning", "Advanced Novice A",
                   "Advanced Novice"]
        self.assertEqual(sorted(classes, key=class_sort_key),
                         ["Oppvisning", "Cubs u/13", "Cubs", "Advanced Novice", "Advanced Novice A",
                          "Junior Kvinner", "Ukjent", ""])

    def test_sort_by_class_keeps_order_within_class_and_puts_short_program_first(self):
        app = App.__new__(App)
        app.refresh_schedule = Mock()
        rows = [self.runner("A", Klasse="Basic", Segment="FS"), self.runner("B", Klasse="Oppvisning"),
                self.runner("C", Klasse="Basic", Segment="SP"), self.runner("D", Klasse="Oppvisning"),
                self.runner("E")]
        app.items = layout_items(rows, 1, 8, 0)

        app.sort_by_class()

        self.assertEqual([r["GivenName"] for r in app.rows], ["B", "D", "C", "A", "E"])

    def test_startliste_has_class_for_runs_only(self):
        items = layout_items([self.runner("A", Klasse="Springs")], 1, 8, 0)
        entries = build_startliste(items, 190, datetime(2000, 1, 1, 8, 0))
        self.assertEqual([e["klasse"] for e in entries], ["", "Springs"])

    def test_startliste_has_program_column_instead_of_suffix(self):
        row = self.runner("A", Ovelse="Kortprogram+Friløp")
        items = layout_items(split_segments([row]), 1, 8, 0)
        entries = build_startliste(items, 190, datetime(2000, 1, 1, 8, 0))
        self.assertEqual([(e["navn"], e["program"]) for e in entries[1:]],
                         [("A Test", "Kortprogram"), ("A Test", "Friløp")])

    def test_pdf_columns_fill_width_and_only_text_columns_shrink(self):
        self.assertEqual(fit_columns([20, 50, 100], [2], 200), [20, 50, 130])
        widths = fit_columns([20, 50, 100, 100], [2, 3], 220)
        self.assertEqual(widths[:2], [20, 50])
        self.assertAlmostEqual(sum(widths), 220)

    def test_startliste_pdf_is_written_without_club_column(self):
        items = layout_items([self.runner("Alejandra Bailador", "Hognestad", Klasse="Oppvisning-cubs")], 1, 8, 0)
        entries = build_startliste(items, 190, datetime(2000, 1, 1, 8, 0))
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "startliste.pdf"
            self.assertTrue(generate_startliste_pdf(entries, str(path), "Tittel", Mock()))
            self.assertTrue(path.read_bytes().startswith(b"%PDF"))

    def test_fp_in_filename_is_free_skate(self):
        self.assertEqual(parse_music_path("Andjelija_Grabyte_FP.mp3")["segment"], "FS")
        self.assertEqual(parse_music_path("Frida_Qianlu_He_SP.mp3")["segment"], "SP")

    def test_sorting_keeps_cancelled_runners_out_of_active_slots(self):
        app = App.__new__(App)
        app.refresh_schedule = Mock()
        rows = [self.runner("C"), self.runner("A"), self.runner("B", Påmelding="Avmeldt")]
        app.items = layout_items(rows, 1, 8, 0)

        app.sort_by_given()

        self.assertEqual([r["GivenName"] for r in app.rows], ["A", "C", "B"])

    def test_validates_setup_before_loading(self):
        valid = {"type": "kunstlopassistenten-oppsett", "version": 1, "settings": {"start": "18:00"},
                 "items": [{"Kind": "warmup"}, self.runner("A", LopLengde=150)]}
        self.assertIsNone(validate_setup(valid))
        for broken in (
            dict(valid, type="noe annet"),
            dict(valid, version=2),
            dict(valid, settings=["feil"]),
            dict(valid, items=["ikke en rad"]),
            dict(valid, items=[self.runner("A", LopLengde="3:00")]),
            dict(valid, items=[self.runner("A", LopLengde=-60)]),
            dict(valid, items=[self.runner("A", Påmelding=5)]),
        ):
            self.assertIsNotNone(validate_setup(broken), broken)

    def test_reads_only_attending_from_spond(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "spond.xlsx"
            wb = openpyxl.Workbook()
            for row in (["Prøvestevne"], ["lør. 31. oktober 08:15 - 11:45"], ["Iskanten"], [None],
                        ["Deltar (2)"], ["Navn", "Foresatt"], ["Anna  Marie Berg"], [None, "Forelder"],
                        ["Ola Hansen"], ["Ikke svart (1)"], ["Navn"], ["Kari Nordmann"]):
                wb.active.append(row)
            wb.save(path)

            rows, info = load_participants_from_spond(path, Mock())

        self.assertEqual([(r["GivenName"], r["FamilyName"]) for r in rows],
                         [("Anna Marie", "Berg"), ("Ola", "Hansen")])
        self.assertEqual((info["title"], info["start"], info["location"]), ("Prøvestevne", "08:15", "Iskanten"))
        self.assertEqual(info["date"].strftime("%d.%m"), "31.10")

    def test_spond_date_picks_nearest_year(self):
        today = datetime(2026, 12, 20).date()
        self.assertEqual(parse_spond_date("lør. 9. januar 10:00", today),
                         (datetime(2027, 1, 9).date(), "10:00"))

    def test_program_comes_from_music_files(self):
        index = {f: parse_music_path(f) for f in ("Bjørg_Mariann_Bjarndisson_FS.mp3",
                                                   "Alva_Teodora_Bjarndisson_SP.mp3",
                                                   "Alva_Teodora_Bjarndisson_FS.mp3")}
        self.assertEqual(ovelse_from_music(self.runner("Bjørg Mariann", "Bjarndisson"), index), "Friløp")
        self.assertEqual(ovelse_from_music(self.runner("Alva Teodora", "Bjarndisson"), index),
                         "Kortprogram+Friløp")
        self.assertEqual(ovelse_from_music(self.runner("Ida", "Chen"), index), "")

    def test_manual_music_is_kept_when_matching(self):
        app = App.__new__(App)
        app.music_files = ["/m/Anna_Test_FS.mp3"]
        app.music_index = {f: parse_music_path(Path(f).name) for f in app.music_files}
        app.music_file_info, app.music_durations, app.music_status = {}, {}, {}
        row = self.runner("Anna", Segment="FS", MusikkFil="/egen/sang.mp3", MusikkManuell=True)

        app.match_rows([row], set())

        self.assertEqual(row["MusikkFil"], "/egen/sang.mp3")

    def test_setup_package_round_trip_with_music(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            song = temp / "Anna_Test_FS.mp3"
            song.write_bytes(b"ID3-disk")
            items = [{"Kind": "warmup"},
                     self.runner("Anna", MusikkFil=str(song), LopLengde=150),
                     self.runner("Ola", MusikkFil="Stevne/FS/Anna_Test_FS.mp3"),
                     self.runner("Kari", MusikkFil=str(temp / "borte.mp3"))]
            data = {"type": "kunstlopassistenten-oppsett", "version": 1, "settings": {"start": "08:15"},
                    "items": items}
            music = {str(song): b"ID3-disk", "Stevne/FS/Anna_Test_FS.mp3": b"ID3-zip"}

            missing = export_setup_package(data, music.get, temp / "pakke.zip")
            imported = extract_setup_package(temp / "pakke.zip", temp / "ny")

            self.assertEqual(missing, ["Kari Test"])
            rows = imported["items"][1:]
            self.assertEqual([Path(r["MusikkFil"]).read_bytes() for r in rows[:2]], [b"ID3-disk", b"ID3-zip"])
            self.assertEqual(rows[2]["MusikkFil"], "")
            self.assertEqual((imported["settings"], rows[0]["LopLengde"]), ({"start": "08:15"}, 150))
            self.assertEqual(items[1]["MusikkFil"], str(song))

    def test_setup_package_keeps_files_that_differ_only_in_case(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            items = [self.runner("A", MusikkFil="x/Sang.mp3"), self.runner("B", MusikkFil="y/sang.mp3")]
            music = {"x/Sang.mp3": b"ID3-a", "y/sang.mp3": b"ID3-b"}

            export_setup_package({"type": "kunstlopassistenten-oppsett", "items": items}, music.get,
                                 temp / "pakke.zip")
            imported = extract_setup_package(temp / "pakke.zip", temp / "ny")

            self.assertEqual([Path(r["MusikkFil"]).read_bytes() for r in imported["items"]],
                             [b"ID3-a", b"ID3-b"])

    def test_import_that_cannot_be_saved_is_not_reported_as_done(self):
        app = App.__new__(App)
        app.confirm_discard = Mock(return_value=True)
        app.apply_setup = Mock()
        app.save_setup = Mock(return_value=False)
        app.update_title, app.log, app.show_folder_link = Mock(), Mock(), Mock()
        with tempfile.TemporaryDirectory() as temp_dir:
            package = Path(temp_dir) / "pakke.zip"
            export_setup_package({"type": "kunstlopassistenten-oppsett", "items": []}, lambda f: None, package)
            with patch("kunstlopassistenten.filedialog.askopenfilename", return_value=str(package)), \
                    patch("kunstlopassistenten.Path.home", return_value=Path(temp_dir)), \
                    patch("kunstlopassistenten.messagebox.showwarning") as warning:
                app.import_package()

        warning.assert_called_once()
        app.show_folder_link.assert_not_called()
        self.assertTrue(app.dirty)

    def test_setup_package_cannot_write_outside_folder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            data = {"type": "kunstlopassistenten-oppsett", "items": [self.runner("A", MusikkFil="musikk/../x.mp3")]}
            with zipfile.ZipFile(temp / "pakke.zip", "w") as package:
                package.writestr("oppsett.json", json.dumps(data))
                package.writestr("x.mp3", b"ID3")

            imported = extract_setup_package(temp / "pakke.zip", temp / "ny")

            self.assertEqual(imported["items"][0]["MusikkFil"], "")
            self.assertFalse((temp / "ny" / "x.mp3").exists())

    def test_playlist_uses_updated_music_from_zip(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            music_zip = Path(temp_dir) / "Musikk.zip"
            out_dir = Path(temp_dir) / "out"
            rows = [{"MusikkFil": "mappe/sang.mp3", "PrintName": "Test"}]
            for data in (b"ID3-gammel", b"ID3-ny"):
                with zipfile.ZipFile(music_zip, "w") as archive:
                    archive.writestr("mappe/sang.mp3", data)
                self.assertTrue(generate_vlc_playlist(rows, out_dir, "x", music_zip, Mock()))

            self.assertEqual((out_dir / "music" / "mappe" / "sang.mp3").read_bytes(), b"ID3-ny")

    def test_new_playback_stops_running_vlc(self):
        app = App.__new__(App)
        old_process = Mock()
        old_process.poll.return_value = None
        app.vlc_process = old_process
        app.audio_ready = False
        app.external_playback = True

        app.stop_running_playback()

        old_process.terminate.assert_called_once()
        self.assertIsNone(app.vlc_process)
        self.assertFalse(app.external_playback)


class PlayerTests(unittest.TestCase):
    def make_app(self, paused=False):
        # Avspiller med falsk mixer: ett spor på 120 s som spilles internt.
        app = App.__new__(App)
        app.music = Mock()
        app.music.get_pos.return_value = 5000
        app.music.get_busy.return_value = not paused
        app.audio_backend = SimpleNamespace(mixer=SimpleNamespace(music=app.music))
        app.audio_ready = True
        app.vlc_process = None
        app.external_playback = False
        app.player_enabled = True
        app.seek_supported = True
        app.use_external_player_var = FakeVariable(False)
        app.current_track = "sang.mp3"
        app.current_duration = 120
        app.is_paused = paused
        app.play_offset = 0.0
        app.drag_position = None
        app.preview_track = ""
        app.wave_peaks = [0.5]
        app.wave_progress = 0.0
        app.wave_cache = {}
        app.wave_path = None
        app.wave_loading = set()
        app.wave_results = queue.Queue()
        for name in ("play_pause_text", "player_track_var", "player_time_var", "player_total_var",
                     "status_var", "music_info_var"):
            setattr(app, name, FakeVariable(""))
        app.player_buttons = {}
        app.log = Mock()
        app.draw_waveform = Mock()
        app.set_wave_progress = Mock()
        app.wave_canvas = SimpleNamespace(winfo_width=lambda: 100)
        app.root = Mock()
        return app

    def test_vlc_toggle_stops_internal_playback(self):
        app = self.make_app()
        app.use_external_player_var.set(True)

        app.on_use_external_player_toggle()

        app.music.stop.assert_called()
        self.assertFalse(app.player_enabled)
        self.assertEqual(app.current_track, "")
        self.assertEqual(app.player_time_var.get(), "VLC direkte")

    def test_vlc_text_is_not_overwritten_by_old_time(self):
        app = self.make_app()
        app.use_external_player_var.set(True)
        app.on_use_external_player_toggle()
        app.current_duration = 120

        app.update_player_ui()

        self.assertEqual(app.player_time_var.get(), "VLC direkte")

    def test_new_setup_resets_player(self):
        app = self.make_app()

        app.reset_music()

        app.music.stop.assert_called()
        self.assertEqual((app.current_track, app.current_duration, app.wave_peaks), ("", 0, []))
        self.assertEqual(app.player_track_var.get(), "Ingen spor valgt")

    def test_seek_while_paused_keeps_audio_paused(self):
        app = self.make_app(paused=True)

        app.seek_to(30)

        app.music.set_pos.assert_called_once_with(30)
        app.music.play.assert_not_called()
        app.music.unpause.assert_not_called()
        self.assertEqual(app.current_position(), 30)
        self.assertEqual(app.player_time_var.get(), "0:30")

        app.toggle_play_pause()

        app.music.unpause.assert_called_once()
        app.music.play.assert_not_called()

    def test_failed_seek_keeps_playing_and_disables_seeking(self):
        app = self.make_app()
        app.play_offset = 40
        app.music.set_pos.side_effect = Exception("ikke støttet")

        app.seek_to(30)

        app.music.play.assert_not_called()
        app.music.stop.assert_not_called()
        self.assertEqual(app.current_position(), 45)
        self.assertFalse(app.can_seek())
        self.assertIn("Spoling støttes ikke", app.status_var.get())

        app.skip_seconds(10)

        app.music.set_pos.assert_called_once()

    def test_drag_seeks_once_on_release(self):
        app = self.make_app()

        for x in (10, 20, 50):
            app.on_waveform_drag(SimpleNamespace(x=x))
        app.music.set_pos.assert_not_called()
        self.assertEqual(app.player_time_var.get(), "1:00")

        app.on_waveform_release(None)

        app.music.set_pos.assert_called_once_with(60)
        app.music.play.assert_not_called()

    def test_skip_uses_current_position(self):
        app = self.make_app()
        app.play_offset = 40

        app.skip_seconds(-10)

        app.music.set_pos.assert_called_once_with(35)
        self.assertEqual(app.current_position(), 35)

    def make_preview_app(self):
        app = self.make_app()
        app.current_track = ""
        app.current_duration = 0
        app.items = [{"MusikkFil": "sang.mp3", "MusikkSek": "95", "NavnFraIsonen": "Anna Berg"}]
        app.get_selected_mp3_filename = Mock(return_value="sang.mp3")
        app.get_cached_mp3_path = Mock(return_value=Path("sang.mp3"))
        app.ensure_audio_backend = Mock(return_value=True)
        app.wave_cache[str(Path("sang.mp3"))] = ([0.2, 1.0], 96.0)
        return app

    def test_selecting_row_shows_waveform(self):
        app = self.make_preview_app()

        app.preview_selected_waveform()

        app.get_cached_mp3_path.assert_called_once_with("sang.mp3", quiet=True)
        self.assertEqual(app.wave_peaks, [0.2, 1.0])
        self.assertEqual(app.player_total_var.get(), "1:35")
        self.assertEqual(app.player_track_var.get(), "Anna Berg - sang.mp3")
        app.music.play.assert_not_called()

    def test_selecting_row_while_playing_keeps_current_track(self):
        app = self.make_preview_app()
        app.current_track = "annen.mp3"

        app.preview_selected_waveform()

        app.get_cached_mp3_path.assert_not_called()

    def test_waveform_is_computed_in_background(self):
        app = self.make_preview_app()
        app.wave_cache = {}

        with patch("kunstlopassistenten.threading.Thread") as thread:
            app.preview_selected_waveform()

        thread.return_value.start.assert_called_once()
        self.assertEqual(app.wave_peaks, [])
        self.assertEqual(app.player_total_var.get(), "1:35")

        app.wave_results.put((str(Path("sang.mp3")), [0.3, 1.0], 96.0, ""))
        app.poll_waveforms()

        self.assertEqual(app.wave_peaks, [0.3, 1.0])
        self.assertEqual(app.wave_loading, set())

    def test_old_waveform_result_is_ignored(self):
        app = self.make_preview_app()
        app.wave_path = "ny.mp3"
        app.wave_peaks = []
        app.wave_loading = {"gammel.mp3"}

        app.wave_results.put(("gammel.mp3", [1.0], 50.0, ""))
        app.poll_waveforms()

        self.assertEqual(app.wave_peaks, [])
        self.assertIn("gammel.mp3", app.wave_cache)

    def test_compute_waveform_normalizes_peaks(self):
        samples = array("h", [0, 100, -200, 50] * 400)
        sound = SimpleNamespace(get_length=lambda: 2.0, get_raw=lambda: samples.tobytes())
        mixer = SimpleNamespace(Sound=lambda path: sound)

        peaks, length, error = compute_waveform(mixer, "sang.mp3")

        self.assertEqual((length, error), (2.0, ""))
        self.assertEqual(max(peaks), 1.0)

    def test_click_in_preview_starts_once_from_position(self):
        app = self.make_preview_app()
        app.preview_selected_waveform()

        app.on_waveform_drag(SimpleNamespace(x=50))
        app.music.play.assert_not_called()
        app.on_waveform_release(None)

        self.assertEqual(app.music.play.call_args_list, [call(start=47.5)])
        self.assertEqual(app.current_track, "sang.mp3")
        self.assertEqual(app.play_offset, 47.5)

    def test_unsupported_start_position_does_not_start_from_beginning(self):
        app = self.make_preview_app()
        app.preview_selected_waveform()
        app.music.play.side_effect = Exception("ikke støttet")
        app.start_vlc_playback = Mock()

        app.on_waveform_drag(SimpleNamespace(x=50))
        app.on_waveform_release(None)

        self.assertEqual(app.music.play.call_args_list, [call(start=47.5)])
        app.start_vlc_playback.assert_not_called()
        self.assertEqual(app.current_track, "")
        self.assertIn("Spoling støttes ikke", app.status_var.get())

    def test_vlc_fallback_gets_start_position(self):
        app = self.make_preview_app()
        app.preview_selected_waveform()
        app.music.load.side_effect = Exception("ukjent format")

        with patch("kunstlopassistenten.find_vlc_executable", return_value="vlc.exe"), \
                patch("kunstlopassistenten.subprocess.Popen") as popen:
            app.on_waveform_drag(SimpleNamespace(x=50))
            app.on_waveform_release(None)

        self.assertIn("--start-time=47.5", popen.call_args.args[0])
        app.music.play.assert_not_called()

    def test_without_vlc_position_does_not_start_default_player_from_beginning(self):
        app = self.make_preview_app()
        app.preview_selected_waveform()
        app.music.load.side_effect = Exception("ukjent format")

        with patch("kunstlopassistenten.find_vlc_executable", return_value=None), \
                patch("kunstlopassistenten.os.startfile", create=True) as startfile:
            app.on_waveform_drag(SimpleNamespace(x=50))
            app.on_waveform_release(None)

        startfile.assert_not_called()

    def test_stopped_track_can_be_restarted_from_waveform(self):
        app = self.make_app()
        app.ensure_audio_backend = Mock(return_value=True)

        app.stop_playback()

        self.assertEqual(app.preview_track, "sang.mp3")
        self.assertTrue(app.can_start_from_wave())

    def test_quiet_cache_lookup_shows_no_dialogs(self):
        app = App.__new__(App)
        app.music_zip = None

        with patch("kunstlopassistenten.messagebox") as messagebox:
            self.assertIsNone(app.get_cached_mp3_path("sang.mp3", quiet=True))
            self.assertIsNone(app.get_cached_mp3_path("", quiet=True))

        messagebox.showerror.assert_not_called()
        messagebox.showinfo.assert_not_called()

    def test_disabled_skip_button_ignores_keyboard(self):
        app = self.make_app()
        app.seek_supported = False
        command = Mock()

        self.assertEqual(app.press_player_button("forward", command), "break")
        command.assert_not_called()
        app.press_player_button("stop", command)
        command.assert_called_once()


if __name__ == "__main__":
    unittest.main()
