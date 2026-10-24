import os
import sys
import zipfile
import csv
import unicodedata
from pathlib import Path, PurePosixPath
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, simpledialog
from tkinter.scrolledtext import ScrolledText
from datetime import datetime, timedelta
import subprocess
import tempfile
from io import BytesIO
import hashlib
import random
import re
import json
import shutil
import wave
from tkinter import font as tkfont
from xml.sax.saxutils import escape as xml_escape

APP_NAME = "Kunstløpassistenten"

# Farger: mørk skifer med isblå aksent (samme oppsett som pc_debugger, annen farge)
BG = "#0F172A"            # vindusbakgrunn
PANEL = "#1B2336"         # paneler
FIELD = "#0B1222"         # felt, tabeller og logg
LINE = "#334155"          # kanter
TEXT = "#F8FAFC"          # hovedtekst
TEXT_DIM = "#94A3B8"      # sekundær tekst
TEXT_DISABLED = "#5B6B82"
BUTTON_HOVER = "#243049"
ACCENT = "#38BDF8"        # hovedknapper og markering
ACCENT_HOVER = "#7DD3FC"
ACCENT_TEXT = "#0F172A"
GROUP_BG = "#26324A"      # oppvarmingsgrupper i tidsplanen
WARN = "#F2B33D"
GOOD = "#86D9A3"
BAD = "#F87171"

SETUP_FILETYPES = [("Oppsett", "*.json"), ("Alle filer", "*.*")]

# Programmene i musikktjenesten: mappenavn/filendelse i musikk-zip -> navn.
SEGMENT_NAMES = {"SP": "Kortprogram", "FS": "Friløp"}
# Status fra music-status.csv i musikk-zip.
MUSIC_STATUS_TEXT = {
    "included": "Mangler i zip",
    "music missing": "Ikke lastet opp",
    "profile not matched": "Profil ikke koblet",
}
AUDIO_SUFFIXES = {".mp3", ".wav", ".m4a", ".mp4", ".aac"}
# Avspilleren: spoling med « » og antall streker i lydbølgen.
SKIP_SECONDS = 10
WAVE_BARS = 160
# Kolonner i deltakerrapporten (Excel, HTML og PDF).
REPORT_HEADERS = ["PrintName", "Organisation", "Klasse", "Program", "Påmelding", "StartTid", "SluttTid",
                  "MusikkStatus", "MusikkTid"]


def safe_int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def is_cancelled(status):
    return "avmeld" in (status or "").strip().lower()


def normalize_text(value):
    text = str(value or "").strip().lower()
    text = (
        text.replace("ø", "o")
        .replace("å", "a")
        .replace("æ", "ae")
        .replace("ö", "o")
        .replace("ä", "a")
        .replace("é", "e")
    )
    # Øvrige aksenter (ü, ñ, ...) fjernes, så navn i filnavn og lister blir like.
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def tokenize_name(value):
    norm = normalize_text(value)
    return [t for t in norm.split() if t]


def full_name_key(row):
    return normalize_text(f"{row.get('GivenName', '')} {row.get('FamilyName', '')}")


def segments_for(ovelse):
    # «Kortprogram+Friløp» -> ["SP", "FS"]. Alt annet regnes som ett friløp.
    text = normalize_text(ovelse)
    segments = ["SP"] if "kortprogram" in text else []
    if "frilop" in text or not segments:
        segments.append("FS")
    return segments


def split_segments(rows):
    # Utøvere med både kortprogram og friløp får én rad per program. Kortprogrammet
    # beholder plassen i listen, friløpene legges bakerst i samme rekkefølge.
    out, free_skates = [], []
    for row in rows:
        if row.get("Segment"):
            out.append(row)
            continue
        segments = segments_for(row.get("Ovelse"))
        row["Segment"] = segments[0]
        out.append(row)
        if len(segments) > 1:
            row["Program"] = SEGMENT_NAMES["SP"]
            copy = dict(row, Segment="FS", Program=SEGMENT_NAMES["FS"], MusikkFil="", Musikk="",
                        MusikkStatus="", MusikkTid="", MusikkSek="", LopGjennomfort=False)
            copy.pop("LopLengde", None)
            free_skates.append(copy)
    return out + free_skates


def add_missing_segments(items):
    # Utøvere som får «Kortprogram+Friløp» etter at listen er delt opp (f.eks. ved senere import
    # fra musikktjenesten): kortprogrammet tar plassen, friløpsraden flyttes bakerst.
    moved = []
    for k, item in enumerate(items):
        if is_marker(item) or item.get("Program") or len(segments_for(item.get("Ovelse"))) < 2:
            continue
        items[k] = dict(item, Segment="SP", Program=SEGMENT_NAMES["SP"], MusikkFil="", Musikk="",
                        MusikkStatus="", MusikkTid="", MusikkSek="", LopGjennomfort=False)
        items[k].pop("LopLengde", None)
        item["Segment"], item["Program"] = "FS", SEGMENT_NAMES["FS"]
        moved.append(item)
    items.extend(moved)
    return len(moved)


def parse_music_path(filename):
    # Musikktjenesten: <stevne>/<Klasse_Kjønn>/<SP|FS>/<Fornavn_Etternavn>_<SP|FS>.mp3
    path = PurePosixPath(filename)
    match = re.match(r"(.+)_(SP|FS)$", path.stem, re.IGNORECASE)
    if not match:
        return None
    parts = path.parts
    return {
        "name": normalize_text(match.group(1).replace("_", " ")),
        "segment": match.group(2).upper(),
        "category": normalize_text(parts[-3].replace("_", " ")) if len(parts) >= 3 else "",
    }


def parse_music_status_csv(text):
    # music-status.csv: Name,Club,Category,Segment,Status -> {(navn, program): status}
    status = {}
    for line in csv.DictReader(text.splitlines()):
        key = (normalize_text(line.get("Name")), (line.get("Segment") or "").strip().upper())
        status[key] = (line.get("Status") or "").strip().lower()
    return status


def find_music_for_row(row, music_index, used_files, fuzzy=True):
    # Finner lydfilen for utøverens program (SP/FS) i en strukturert musikk-zip.
    # Først eksakt navn, så (fuzzy) fornavn + alle etternavn. Ved like navn vinner riktig klasse.
    # Returnerer (fil, usikker); usikker betyr at flere filer passer og koblingen bør sjekkes.
    segment = row.get("Segment") or "FS"
    files = [f for f, info in music_index.items() if info and info["segment"] == segment]
    name = full_name_key(row)
    hits = [f for f in files if music_index[f]["name"] == name]
    if not hits and fuzzy:
        tokens = tokenize_name(row.get("GivenName"))[:1] + tokenize_name(row.get("FamilyName"))
        if len(tokens) > 1:
            hits = [f for f in files if set(tokens) <= set(music_index[f]["name"].split())]
    klasse = normalize_text(row.get("Klasse"))
    in_class = [f for f in hits if klasse and music_index[f]["category"].startswith(klasse)]
    pool = in_class if any(f not in used_files for f in in_class) else hits
    free = [f for f in pool if f not in used_files]
    if not free:
        return "", False
    used_files.add(free[0])
    return free[0], len(pool) > 1


def music_status_text(row, matched, status_map, uncertain=False):
    if matched:
        return "Usikker kobling – sjekk lydfil" if uncertain else "OK"
    status = status_map.get((full_name_key(row), row.get("Segment") or "FS"))
    if status is None:
        return "Ikke i musikktjenesten" if status_map else "Mangler musikk"
    return MUSIC_STATUS_TEXT.get(status, status)


def sanitize_filename(value):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip())
    return cleaned.strip("._") or "fil"


def detect_audio_format(data):
    if data[:4] == b"RIFF" and data[8:12] == b"WAVE":
        return "WAV", ".wav"
    if len(data) >= 12 and data[4:8] == b"ftyp":
        return "M4A", ".m4a"
    if data[:3] == b"ID3":
        return "MP3", ".mp3"
    if len(data) >= 2 and data[0] == 0xFF and (data[1] & 0xE0) == 0xE0:
        version_bits = (data[1] >> 3) & 0x03
        layer_bits = (data[1] >> 1) & 0x03
        if version_bits != 0x01 and layer_bits != 0x00:
            return "MP3", ".mp3"
    return "Ukjent", ""


def get_audio_duration(data, audio_format):
    try:
        if audio_format == "WAV":
            with wave.open(BytesIO(data), "rb") as audio:
                frame_rate = audio.getframerate()
                return audio.getnframes() / frame_rate if frame_rate else None
        if audio_format == "MP3":
            from mutagen.mp3 import MP3

            return MP3(BytesIO(data)).info.length
        if audio_format == "M4A":
            from mutagen.mp4 import MP4

            return MP4(BytesIO(data)).info.length
    except Exception:
        return None
    return None


def find_vlc_executable():
    candidates = [shutil.which("vlc")]
    for env_name in ("ProgramFiles", "ProgramFiles(x86)"):
        base_dir = os.environ.get(env_name)
        if base_dir:
            candidates.append(Path(base_dir) / "VideoLAN" / "VLC" / "vlc.exe")
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)
    return None


def get_version():
    version = "ukjent"
    try:
        result = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            capture_output=True,
            text=True,
            cwd=Path(__file__).resolve().parent,
        )
        if result.returncode == 0:
            version = result.stdout.strip() or "ukjent"
    except Exception:
        version = "ukjent"
    if version == "ukjent":
        try:
            version_file = Path(__file__).resolve().parent / "version.txt"
            if version_file.exists():
                version = version_file.read_text(encoding="utf-8").strip() or "ukjent"
        except Exception:
            pass
    if version == "ukjent":
        try:
            mtime = Path(__file__).resolve().stat().st_mtime
            version = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d")
        except Exception:
            pass
    return version


def format_generated_ts():
    computer_name = os.environ.get("COMPUTERNAME") or os.environ.get("HOSTNAME") or ""
    suffix = f" • {computer_name}" if computer_name else ""
    version = get_version()
    return (
        datetime.now().strftime("Generert %d.%m.%Y %H:%M")
        + suffix
        + f" • Revisjon: {version}"
    )


def format_duration(seconds):
    if seconds is None:
        return ""
    total = int(round(seconds))
    mins = total // 60
    secs = total % 60
    return f"{mins}:{secs:02d}"


def format_hms(seconds):
    total = int(seconds)
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def parse_time_hhmm(value):
    if not value:
        return None
    parts = [p.strip() for p in value.split(":")]
    try:
        if len(parts) == 2:
            h, m = int(parts[0]), int(parts[1])
            return datetime(2000, 1, 1, h, m, 0)
        if len(parts) == 3:
            h, m, s = int(parts[0]), int(parts[1]), int(parts[2])
            return datetime(2000, 1, 1, h, m, s)
    except ValueError:
        return None
    return None


def parse_duration_mmss(value):
    if not value:
        return None
    parts = [p.strip() for p in value.split(":")]
    try:
        if len(parts) == 2:
            h, m, s = 0, int(parts[0]), int(parts[1])
        elif len(parts) == 3:
            h, m, s = int(parts[0]), int(parts[1]), int(parts[2])
        else:
            return None
    except ValueError:
        return None
    # Negative tider og sekunder (eller minutter i T:MM:SS) over 59 avvises.
    if min(h, m, s) < 0 or s > 59 or (len(parts) == 3 and m > 59):
        return None
    return h * 3600 + m * 60 + s


def parse_date_ddmmyy(value):
    try:
        return datetime.strptime(value.strip(), "%d.%m.%y").date()
    except Exception:
        return None


def format_date_long(date_obj):
    months = [
        "januar",
        "februar",
        "mars",
        "april",
        "mai",
        "juni",
        "juli",
        "august",
        "september",
        "oktober",
        "november",
        "desember",
    ]
    return f"{date_obj.day} {months[date_obj.month - 1]} {date_obj.year}"


def is_registered(status):
    if is_cancelled(status):
        return False
    if not status:
        return False
    text = str(status).strip().lower()
    if "ikke sjekket inn" in text:
        return True
    return "påmeld" in text or "registr" in text or "bekreftet" in text


def new_participant(given, family, club="", gender="", status=""):
    given = str(given or "").strip()
    family = str(family or "").strip()
    return {
        "PrintName": f"{family}, {given}".strip(", "),
        "NavnFraIsonen": f"{given} {family}".strip(),
        "GivenName": given,
        "FamilyName": family,
        "Gender": str(gender or "").strip(),
        "Organisation": str(club or "").strip(),
        "Påmelding": str(status or "").strip(),
        "NifId": "",
        "Klasse": "",
        "Ovelse": "",
        # Segment: SP (kortprogram) eller FS (friløp). Program settes bare når utøveren har begge.
        "Segment": "",
        "Program": "",
        "Musikk": "",
        "MusikkStatus": "",
        "MusikkFil": "",
        "MusikkFormat": "",
        "MusikkFormatAvvik": False,
        "MusikkSpillbar": False,
        "MusikkTid": "",
        "MusikkSek": "",
        "StartTid": "",
        "SluttTid": "",
    }


def split_name(full_name):
    parts = (full_name or "").split()
    if len(parts) < 2:
        return " ".join(parts), ""
    return " ".join(parts[:-1]), parts[-1]


def display_name(row):
    full = f"{row.get('GivenName', '')} {row.get('FamilyName', '')}".strip()
    return full or row.get("NavnFraIsonen") or row.get("PrintName") or ""


def run_name(row):
    # Navn i tidsplanen; utøvere med to program får programmet bak navnet.
    program = row.get("Program")
    return f"{display_name(row)} – {program}" if program else display_name(row)


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


SETUP_TEXT_FIELDS = ("PrintName", "NavnFraIsonen", "GivenName", "FamilyName", "Organisation", "Påmelding",
                     "Klasse", "Ovelse", "Segment", "Program", "MusikkFil")


def validate_setup(data):
    # Sjekker hele oppsettfilen før noe tas i bruk. Returnerer feilmelding, eller None når den er gyldig.
    if data.get("type") not in (None, "kunstlopassistenten-oppsett"):
        return "Filen er ikke et oppsett fra Kunstløpassistenten."
    if safe_int(data.get("version"), 1) > 1:
        return "Oppsettet er laget med en nyere versjon av Kunstløpassistenten."
    settings = data.get("settings") or {}
    if not isinstance(settings, dict) or not all(isinstance(v, str) for v in settings.values()):
        return "Innstillingene i oppsettet er ugyldige."
    for nr, item in enumerate(data["items"], start=1):
        if not isinstance(item, dict):
            return f"Rad {nr} i oppsettet er ugyldig."
        if item.get("Kind") not in (None, "warmup", "pause"):
            return f"Rad {nr} har ukjent type: {item.get('Kind')}"
        for field in SETUP_TEXT_FIELDS:
            if item.get(field) is not None and not isinstance(item[field], str):
                return f"Rad {nr} har ugyldig verdi i feltet {field}."
        length = item.get("LopLengde")
        if length not in (None, "") and (isinstance(length, bool) or not isinstance(length, int) or length <= 0):
            return f"Rad {nr} ({display_name(item)}) har ugyldig lengde på løp: {length}"
    return None


def load_participants_from_excel(excel_path, log):
    try:
        import openpyxl
    except Exception:
        log("Mangler openpyxl. Installer med: pip install openpyxl")
        return []

    if not Path(excel_path).exists():
        log(f"Fant ikke excel: {excel_path}")
        return []

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
    except Exception as exc:
        log(f"Kunne ikke lese excel: {excel_path} ({exc})")
        return []

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        log("Excel er tom.")
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    header_map = {h: idx for idx, h in enumerate(headers)}

    def idx(*names):
        for name in names:
            if name in header_map:
                return header_map[name]
        return None

    i_given = idx("Fornavn")
    i_family = idx("Etternavn")
    i_gender = idx("Kjønn")
    i_club = idx("Klubb")
    i_status = idx("Påmelding", "Påmeldingsstatus", "Deltakerstatus")
    # Person ID er samme nummer som NIF ID i musikktjenesten.
    extra_fields = {"NifId": idx("Person ID"), "Klasse": idx("Klasse"), "Ovelse": idx("Øvelse")}

    if i_given is None or i_family is None:
        log("Finner ikke nødvendige kolonner i excel (Fornavn/Etternavn).")
        return []

    out = []
    for row in rows[1:]:
        given = row[i_given] if i_given < len(row) else ""
        family = row[i_family] if i_family < len(row) else ""
        gender = row[i_gender] if i_gender is not None and i_gender < len(row) else ""
        club = row[i_club] if i_club is not None and i_club < len(row) else ""
        status = row[i_status] if i_status is not None and i_status < len(row) else ""

        if not (given or family):
            continue

        participant = new_participant(given, family, club, gender, status)
        for field, i in extra_fields.items():
            if i is not None and i < len(row):
                participant[field] = cell_text(row[i])
        out.append(participant)

    log(f"Lest excel: {excel_path} ({len(out)} rader)")
    return out


def load_entries_from_excel(excel_path, log):
    # Deltakereksport fra musikktjenesten (entries_*.xlsx). Over tabellen står stevnenavn og
    # «Arrangør · DD.MM.ÅÅÅÅ». Returnerer (deltakere, {"title": ..., "date": date}).
    try:
        import openpyxl
    except Exception:
        log("Mangler openpyxl. Installer med: pip install openpyxl")
        return [], {}

    try:
        ws = openpyxl.load_workbook(excel_path).worksheets[0]
    except Exception as exc:
        log(f"Kunne ikke lese excel: {excel_path} ({exc})")
        return [], {}

    rows = list(ws.iter_rows(values_only=True))
    header_row = next(
        (i for i, row in enumerate(rows[:20])
         if {"First name", "Last name"} <= {cell_text(c) for c in row}),
        None,
    )
    if header_row is None:
        log("Finner ikke kolonnene First name/Last name. Er dette deltakerlisten fra musikktjenesten?")
        return [], {}

    header_map = {cell_text(h): i for i, h in enumerate(rows[header_row]) if cell_text(h)}
    columns = {
        "GivenName": "First name",
        "FamilyName": "Last name",
        "Fodselsdato": "Date of birth",
        "Gender": "Gender",
        "Organisation": "Club",
        "Klasse": "Category",
        "Ovelse": "Segments",
        "NifId": "NIF ID",
    }

    def value(row, col):
        i = header_map.get(col)
        return cell_text(row[i]) if i is not None and i < len(row) else ""

    entries = []
    for row in rows[header_row + 1:]:
        entry = {field: value(row, col) for field, col in columns.items()}
        if entry["GivenName"] or entry["FamilyName"]:
            entries.append(entry)

    info = {}
    heading = [cell_text(row[0]) for row in rows[:header_row] if row and cell_text(row[0])]
    if heading:
        info["title"] = heading[0]
    for line in heading[1:]:
        match = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", line)
        if match:
            day, month, year = (int(g) for g in match.groups())
            try:
                info["date"] = datetime(year, month, day).date()
            except ValueError:
                pass
            break
    log(f"Lest deltakere fra musikktjenesten: {excel_path} ({len(entries)} rader)")
    return entries, info


def participant_from_entry(entry):
    row = new_participant(entry["GivenName"], entry["FamilyName"], entry["Organisation"],
                          entry["Gender"], "Påmeldt")
    for field in ("NifId", "Klasse", "Ovelse"):
        row[field] = entry[field]
    return row


def generate_excel(rows, out_path, log):
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font
    except Exception:
        log("Mangler openpyxl. Installer med: pip install openpyxl")
        return False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Participants"

    headers = REPORT_HEADERS
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
        if is_cancelled(row.get("Påmelding", "")):
            from openpyxl.styles import PatternFill

            fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            for cell in ws[ws.max_row]:
                cell.fill = fill
                cell.font = Font(color="7A0B0B")

    for idx, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 18

    try:
        wb.save(out_path)
    except PermissionError:
        log(f"Kunne ikke skrive Excel (filen er trolig åpen): {out_path}")
        return False
    except Exception as exc:
        log(f"Kunne ikke skrive Excel: {out_path} ({exc})")
        return False
    log(f"Excel skrevet: {out_path}")
    return True


def generate_html(rows, out_path, title, log):
    headers = REPORT_HEADERS

    def esc(s):
        return (
            str(s)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    rows_html = []
    for row in rows:
        row_style = ""
        if is_cancelled(row.get("Påmelding", "")):
            row_style = ' style="background:#f8d7da;color:#7a0b0b;"'
        cells = "".join(f"<td>{esc(row.get(h, ''))}</td>" for h in headers)
        rows_html.append(f"<tr{row_style}>{cells}</tr>")

    html = f"""<!doctype html>
<html lang="no">
<head>
  <meta charset="utf-8">
  <title>{esc(title)}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 24px; }}
    table {{ border-collapse: collapse; width: 100%; }}
    th, td {{ border: 1px solid #ccc; padding: 6px 8px; text-align: left; }}
    th {{ background: #f4f4f4; }}
  </style>
</head>
<body>
  <h2>{esc(title)}</h2>
  <table>
    <thead>
      <tr>
        {''.join(f'<th>{esc(h)}</th>' for h in headers)}
      </tr>
    </thead>
    <tbody>
      {''.join(rows_html)}
    </tbody>
  </table>
</body>
</html>
"""

    try:
        Path(out_path).write_text(html, encoding="utf-8")
    except PermissionError:
        log(f"Kunne ikke skrive HTML (filen er trolig åpen): {out_path}")
        return False
    except Exception as exc:
        log(f"Kunne ikke skrive HTML: {out_path} ({exc})")
        return False
    log(f"HTML skrevet: {out_path}")
    return True


def generate_pdf(rows, out_path, title, log):
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            SimpleDocTemplate,
            Table,
            TableStyle,
            Paragraph,
            Spacer,
            Image,
        )
    except Exception:
        log("Mangler reportlab. Installer med: pip install reportlab")
        return False

    headers = REPORT_HEADERS
    data = [headers]
    for row in rows:
        data.append([row.get(h, "") for h in headers])

    doc = SimpleDocTemplate(out_path, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Loddefjord IL Kunstløp", styles["Title"]),
        Spacer(1, 6),
    ]
    logo_path = Path(__file__).resolve().parent / "Lil Logo.jpg"
    if logo_path.exists():
        try:
            logo = Image(str(logo_path))
            logo.drawHeight = 50
            logo.drawWidth = 50 * (logo.imageWidth / logo.imageHeight)
            story.append(logo)
            story.append(Spacer(1, 6))
        except Exception:
            log("Klarte ikke å lese logo-bildet.")
    story.append(Paragraph(xml_escape(title), styles["Heading2"]))
    story.append(Spacer(1, 12))

    table = Table(data, repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    for idx, row in enumerate(rows, start=1):
        if is_cancelled(row.get("Påmelding", "")):
            style_cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#F8D7DA")))
            style_cmds.append(("TEXTCOLOR", (0, idx), (-1, idx), colors.HexColor("#7A0B0B")))
    table.setStyle(TableStyle(style_cmds))
    story.append(table)
    story.append(Spacer(1, 6))
    story.append(Paragraph(format_generated_ts(), styles["Normal"]))
    try:
        doc.build(story)
    except PermissionError:
        log(f"Kunne ikke skrive PDF (filen er trolig åpen): {out_path}")
        return False
    except Exception as exc:
        log(f"Kunne ikke skrive PDF: {out_path} ({exc})")
        return False
    log(f"PDF skrevet: {out_path}")
    return True


def water_pause_positions(group_count, water_count):
    # Fordeler vanningspausene jevnt mellom oppvarmingsgruppene (etter gruppe nr. N).
    positions = {round(group_count * k / (water_count + 1)) for k in range(1, water_count + 1)}
    return sorted(p for p in positions if 0 < p < group_count)


def is_marker(item):
    return item.get("Kind") in ("warmup", "pause")


def place_water_pauses(items, water_count):
    # Fjerner gamle vanningspauser og legger nye jevnt fordelt foran oppvarmingsgruppene,
    # eller mellom utøverne når det ikke er noen oppvarmingsgrupper.
    items = [item for item in items if item.get("Kind") != "pause"]
    anchors = [k for k, item in enumerate(items) if item.get("Kind") == "warmup"]
    if not anchors:
        anchors = [k for k, item in enumerate(items)
                   if not is_marker(item) and not is_cancelled(item.get("Påmelding", ""))]
    for position in reversed(water_pause_positions(len(anchors), water_count)):
        items.insert(anchors[position], {"Kind": "pause"})
    return items


def auto_group_size(runner_count, group_count):
    # Antall utøvere delt likt på gruppene (halvparten, en tredjedel, ...), rundet opp.
    if group_count < 1:
        return 1
    return max(1, -(-runner_count // group_count))


def layout_items(rows, group_count, group_size, water_count):
    # Lager rekkefølgen oppvarmingsgruppe, utøvere, ... Avmeldte legges skjult bakerst.
    # Utøvere som ikke får plass havner i siste gruppe.
    active = [r for r in rows if not is_cancelled(r.get("Påmelding", ""))]
    items = []
    if group_count < 1:
        items.extend(active)
    else:
        group_size = max(1, group_size) if group_count > 1 else len(active)
        for group in range(group_count):
            items.append({"Kind": "warmup"})
            last = group == group_count - 1
            items.extend(active[group * group_size : None if last else (group + 1) * group_size])
    items.extend(r for r in rows if is_cancelled(r.get("Påmelding", "")))
    return place_water_pauses(items, water_count)


def group_overflows(items, group_size):
    # {id(oppvarmingsrad): antall utøvere} for grupper med flere utøvere enn tillatt.
    counts = {}
    current = None
    for item in items:
        if item.get("Kind") == "warmup":
            current = id(item)
            counts[current] = 0
        elif current is not None and not is_marker(item) and not is_cancelled(item.get("Påmelding", "")):
            counts[current] += 1
    return {key: count for key, count in counts.items() if count > group_size}


def build_startliste(
    items,
    run_seconds,
    start_time,
    ready_seconds=0,
    warmup_seconds=0,
    water_seconds=0,
    water_label="Vanningspause",
):
    entries = []
    current = start_time
    group = 0
    nr = 0

    def add(kind, seconds, nr, navn, klubb, row, run_length=None):
        nonlocal current
        end = current + timedelta(seconds=seconds)
        entries.append(
            {
                "is_group": kind != "run",
                "kind": kind,
                "start": current.strftime("%H:%M:%S"),
                "end": end.strftime("%H:%M:%S"),
                "duration": format_hms(seconds),
                "run_length": format_hms(run_length) if run_length is not None else "",
                "custom_length": bool(row.get("LopLengde")),
                "nr": nr,
                "navn": navn,
                "klubb": klubb,
                "row": row,
            }
        )
        current = end

    for item in items:
        kind = item.get("Kind")
        if kind == "warmup":
            group += 1
            add("warmup", warmup_seconds, "", f"Oppvarmingsgruppe {group}", "", item)
        elif kind == "pause":
            add("pause", water_seconds, "", water_label, "", item)
        elif not is_cancelled(item.get("Påmelding", "")):
            nr += 1
            # LopLengde (sekunder) overstyrer standard lengde på løp for én utøver.
            run = item.get("LopLengde") or run_seconds
            add("run", run + ready_seconds, nr, run_name(item), item.get("Organisation", ""), item, run)
    return entries


def generate_startliste_excel(entries, out_path, title, log):
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font, PatternFill
    except Exception:
        log("Mangler openpyxl. Installer med: pip install openpyxl")
        return False

    header_fill = PatternFill(start_color="ADADAD", end_color="ADADAD", fill_type="solid")
    group_fill = PatternFill(start_color="D0D0D0", end_color="D0D0D0", fill_type="solid")
    bold = Font(bold=True)
    generated_ts = format_generated_ts()

    def write_sheet(ws, with_duration):
        ws.append([title])
        ws.append([])
        headers = (["Nr.", "Start", "", "Slutt", "Lengde på løp"]
                   + (["Sammenlagt tidsforbruk"] if with_duration else []) + ["Navn", "Klubb"])
        ws.append(headers)
        for cell in ws[3]:
            cell.font = bold
            cell.fill = header_fill

        group = 0
        for entry in entries:
            start = entry["start"]
            if entry["kind"] == "warmup":
                group += 1
                if group > 1:
                    start = f"ca. {start}"
            duration = [entry["duration"]] if with_duration else []
            ws.append([entry["nr"], start, "-", entry["end"], entry["run_length"]] + duration
                      + [entry["navn"], entry["klubb"]])
            if entry["is_group"]:
                for cell in ws[ws.max_row]:
                    cell.font = bold
                    cell.fill = group_fill

        widths = [6, 10, 3, 10, 14] + ([22] if with_duration else []) + [38, 16]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.append([])
        ws.append([generated_ts])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Startliste"
    write_sheet(ws, with_duration=False)
    write_sheet(wb.create_sheet("Med tidsforbruk"), with_duration=True)

    try:
        wb.save(out_path)
    except PermissionError:
        log(f"Kunne ikke skrive Excel (filen er trolig åpen): {out_path}")
        return False
    except Exception as exc:
        log(f"Kunne ikke skrive Excel: {out_path} ({exc})")
        return False

    log(f"Startliste Excel skrevet: {out_path}")
    return True


def generate_startliste_pdf(entries, out_path, title, log):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception:
        log("Mangler reportlab. Installer med: pip install reportlab")
        return False

    font_name = "Helvetica"
    try:
        calibri_path = Path("C:/Windows/Fonts/calibri.ttf")
        if calibri_path.exists():
            pdfmetrics.registerFont(TTFont("Calibri", str(calibri_path)))
            font_name = "Calibri"
    except Exception:
        font_name = "Helvetica"

    body_style = ParagraphStyle(
        "BodyCell",
        fontName=font_name,
        fontSize=10,
        leading=11,
    )
    # De lange overskriftene brytes over to linjer.
    head_style = ParagraphStyle("HeadCell", fontName=font_name, fontSize=11, leading=12, alignment=1)
    data = [["Nr.", "Start", "", "Slutt", Paragraph("Lengde på løp", head_style),
             Paragraph("Sammenlagt tidsforbruk", head_style), "Navn", "Klubb"]]
    for entry in entries:
        # Paragraph tolker tekst som markup, så & og < i navn må escapes.
        name_cell = Paragraph(xml_escape(entry["navn"]), body_style)
        club_cell = Paragraph(xml_escape(entry["klubb"]), body_style)
        data.append([entry["nr"], entry["start"], "-", entry["end"], entry["run_length"], entry["duration"],
                     name_cell, club_cell])

    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "StartTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=16,
        leading=18,
    )
    generated_ts = format_generated_ts()
    story = [Paragraph(xml_escape(title), title_style), Spacer(1, 8)]

    col_widths = [26, 52, 10, 52, 54, 74, 145, 110]
    table = Table(data, repeatRows=1, colWidths=col_widths)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ADADAD")),
        ("FONTNAME", (0, 0), (-1, 0), font_name),
        ("FONTNAME", (0, 1), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, 0), 11),
        ("FONTSIZE", (0, 1), (-1, -1), 10),
        ("ALIGN", (0, 0), (5, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LINEABOVE", (0, 0), (-1, 0), 0.5, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.black),
        ("LINEBEFORE", (0, 0), (0, -1), 0.5, colors.black),
        ("LINEAFTER", (-1, 0), (-1, -1), 0.5, colors.black),
    ]
    row_idx = 1
    for entry in entries:
        if entry["is_group"]:
            style_cmds.append(
                ("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#D0D0D0"))
            )
            style_cmds.append(("FONTNAME", (0, row_idx), (-1, row_idx), "Helvetica-Bold"))
        row_idx += 1
    table.setStyle(TableStyle(style_cmds))
    story.append(table)
    story.append(Spacer(1, 6))
    story.append(Paragraph(generated_ts, ParagraphStyle("Gen", fontName=font_name, fontSize=9)))

    try:
        doc.build(story)
    except PermissionError:
        log(f"Kunne ikke skrive PDF (filen er trolig åpen): {out_path}")
        return False
    except Exception as exc:
        log(f"Kunne ikke skrive PDF: {out_path} ({exc})")
        return False

    log(f"Startliste PDF skrevet: {out_path}")
    return True


def generate_vlc_playlist(rows, out_dir, base_name, music_zip, log):
    if not music_zip:
        log("Ingen musikk-zip funnet, kan ikke lage spilleliste.")
        return False

    playlist_path = Path(out_dir) / f"Startliste_{base_name}.m3u"
    music_out = Path(out_dir) / "music"
    music_out.mkdir(parents=True, exist_ok=True)

    try:
        with zipfile.ZipFile(music_zip, "r") as mz:
            lines = ["#EXTM3U"]
            for row in rows:
                fname = row.get("MusikkFil", "")
                if not fname:
                    continue
                dest_path = music_out / fname
                dest_path.parent.mkdir(parents=True, exist_ok=True)
                # Pakkes alltid ut på nytt, så en oppdatert zip med samme filnavn ikke gir gammel lyd.
                try:
                    mz.extract(fname, music_out)
                except Exception as exc:
                    if not dest_path.exists():
                        log(f"Kunne ikke pakke ut {fname}: {exc}")
                        continue
                    log(f"ADVARSEL: Kunne ikke oppdatere {fname} ({exc}). Bruker filen som ligger der fra før.")
                duration = row.get("MusikkSek", "")
                performer = row.get("PrintName", "")
                song = Path(fname).stem
                title = f"{performer} - {song}".strip(" -")
                extinf = duration if duration != "" else -1
                lines.append(f"#EXTINF:{extinf},{title}")
                lines.append(str(dest_path.resolve()))
    except Exception as exc:
        log(f"Kunne ikke lage spilleliste: {exc}")
        return False

    try:
        playlist_path.write_text("\n".join(lines), encoding="utf-8")
    except Exception as exc:
        log(f"Kunne ikke skrive spilleliste: {playlist_path} ({exc})")
        return False
    log(f"Spilleliste skrevet: {playlist_path}")
    return True


def colorref(color):
    # "#RRGGBB" -> Windows COLORREF (0x00BBGGRR).
    red, green, blue = int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16)
    return blue << 16 | green << 8 | red


def style_window_frame(win):
    # Tittellinje og vinduskant i appens farger. Windows 11 tar fargene direkte,
    # Windows 10 får bare mørk tittellinje. Andre systemer lar vi være.
    if sys.platform != "win32":
        return
    try:
        import ctypes

        win.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(win.winfo_id())
        attributes = (
            (20, 1),                 # DWMWA_USE_IMMERSIVE_DARK_MODE
            (34, colorref(LINE)),    # DWMWA_BORDER_COLOR
            (35, colorref(BG)),      # DWMWA_CAPTION_COLOR
            (36, colorref(TEXT)),    # DWMWA_TEXT_COLOR
        )
        for attribute, value in attributes:
            value = ctypes.c_int(value)
            ctypes.windll.dwmapi.DwmSetWindowAttribute(
                hwnd, attribute, ctypes.byref(value), ctypes.sizeof(value)
            )
    except Exception:
        pass


class App:
    def __init__(self, root):
        self.root = root
        # Hele tidsplanen i rekkefølge: utøvere (deltaker-dict) og markører for
        # oppvarmingsgrupper og vanningspauser ({"Kind": "warmup"/"pause"}).
        self.items = []
        self.entries = []
        self.setup_path = None
        self.dirty = False
        self._loading = False
        self.drag_iid = None
        self.overflow_tips = {}
        self.tip = None
        self.source_dir = None
        # Deltakerlisten fra musikktjenesten (entries_*.xlsx), slik at rekkefølgen på innlesingen ikke betyr noe.
        self.music_entries = []
        self.music_zip = None
        self.music_cache_dir = None
        self.music_file_info = {}
        self.music_files = []
        self.music_durations = {}
        # {lydfil: {"name", "segment", "category"}} for filer med musikktjenestens filnavn.
        self.music_index = {}
        # {(navn, SP/FS): status} fra music-status.csv.
        self.music_status = {}
        self.audio_backend = None
        self.audio_ready = False
        self.vlc_process = None
        self.current_track = ""
        self.is_paused = False
        self.current_duration = 0
        self.external_playback = False
        # Avspillerens posisjon: get_pos() teller fra siste play(), så spoling legges til her.
        self.play_offset = 0.0
        self.player_enabled = True
        self.wave_peaks = []
        self.wave_cache = {}
        self.wave_progress = 0.0
        self.play_pause_text = tk.StringVar(value="Spill")
        self.use_external_player_var = tk.BooleanVar(value=False)

        self._init_fonts()
        self._init_style()
        self._init_vars()
        self._build_menu()
        self._build()
        self.refresh_schedule()
        self.dirty = False
        self.update_title()
        self.update_player_ui()
        self.update_clock()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    @property
    def rows(self):
        return [item for item in self.items if not is_marker(item)]

    @rows.setter
    def rows(self, value):
        self.items = list(value)

    # -------------------- Oppsett --------------------
    def _init_fonts(self):
        fams = set(tkfont.families())
        base = "Bahnschrift" if "Bahnschrift" in fams else "Segoe UI"
        cond = "Bahnschrift SemiCondensed" if "Bahnschrift SemiCondensed" in fams else base
        semi = "Bahnschrift SemiBold" if "Bahnschrift SemiBold" in fams else base
        mono = "Consolas" if "Consolas" in fams else "Courier New"
        self.fonts = {
            "body": (base, 10),
            "bold": (semi, 10),
            "small": (cond, 10),
            "head": (semi, 13),
            "title": (semi, 16),
            "clock": (semi, 96),
            "mono": (mono, 10),
        }

    def _init_style(self):
        r = self.root
        f = self.fonts
        r.configure(bg=BG)
        r.option_add("*TCombobox*Listbox.background", FIELD)
        r.option_add("*TCombobox*Listbox.foreground", TEXT)
        r.option_add("*TCombobox*Listbox.selectBackground", ACCENT)
        r.option_add("*TCombobox*Listbox.selectForeground", ACCENT_TEXT)
        r.option_add("*TCombobox*Listbox.font", f["body"])

        s = ttk.Style(r)
        s.theme_use("clam")
        s.configure(".", background=BG, foreground=TEXT, font=f["body"], fieldbackground=FIELD,
                    bordercolor=LINE, lightcolor=PANEL, darkcolor=BG, troughcolor=FIELD,
                    selectbackground=ACCENT, selectforeground=ACCENT_TEXT, insertcolor=TEXT,
                    focuscolor=ACCENT)
        s.configure("TFrame", background=BG)
        s.configure("Panel.TFrame", background=PANEL)
        s.configure("TLabel", background=BG, foreground=TEXT)
        s.configure("Title.TLabel", font=f["title"])
        s.configure("Dim.TLabel", foreground=TEXT_DIM, font=f["small"])
        s.configure("Panel.TLabel", background=PANEL)
        s.configure("PanelDim.TLabel", background=PANEL, foreground=TEXT_DIM, font=f["small"])
        s.configure("PanelHead.TLabel", background=PANEL, font=f["head"])
        s.configure("PanelWarn.TLabel", background=PANEL, foreground=WARN)

        s.configure("TButton", background=FIELD, foreground=TEXT, bordercolor=LINE,
                    padding=(12, 6), relief="flat")
        s.map("TButton", background=[("disabled", PANEL), ("active", BUTTON_HOVER)],
              foreground=[("disabled", TEXT_DISABLED)])
        s.configure("Accent.TButton", background=ACCENT, foreground=ACCENT_TEXT, bordercolor=ACCENT,
                    font=f["head"])
        s.map("Accent.TButton", background=[("disabled", "#3B5568"), ("active", ACCENT_HOVER)],
              foreground=[("disabled", PANEL)])

        s.configure("TEntry", foreground=TEXT, fieldbackground=FIELD, padding=5)
        s.configure("TSpinbox", foreground=TEXT, fieldbackground=FIELD, arrowcolor=TEXT, padding=4)
        s.configure("TCombobox", foreground=TEXT, fieldbackground=FIELD, arrowcolor=TEXT, padding=4)
        s.configure("TCheckbutton", background=PANEL, foreground=TEXT, indicatorbackground=FIELD,
                    indicatorforeground=ACCENT)
        s.map("TCheckbutton", background=[("active", PANEL)], indicatorbackground=[("selected", FIELD)])

        s.configure("TNotebook", background=BG, borderwidth=0, tabmargins=(0, 0, 0, 0),
                    bordercolor=BG, lightcolor=BG, darkcolor=BG)
        s.configure("TNotebook.Tab", background=BG, foreground=TEXT_DIM, padding=(16, 7),
                    bordercolor=BG, lightcolor=BG, darkcolor=BG, font=f["body"])
        # clam-temaet gir fanene lyse kanter via lightcolor; de settes til fanens egen farge.
        s.map("TNotebook.Tab", background=[("selected", PANEL)], foreground=[("selected", TEXT)],
              lightcolor=[("selected", PANEL), ("", BG)])

        s.configure("Treeview", background=FIELD, fieldbackground=FIELD, foreground=TEXT,
                    rowheight=26, bordercolor=LINE, font=f["body"])
        s.map("Treeview", background=[("selected", ACCENT)], foreground=[("selected", ACCENT_TEXT)])
        s.configure("Treeview.Heading", background=PANEL, foreground=TEXT_DIM, font=f["small"],
                    relief="flat", bordercolor=LINE, padding=(6, 4))
        s.map("Treeview.Heading", background=[("active", BUTTON_HOVER)])
        s.configure("Horizontal.TProgressbar", background=ACCENT, troughcolor=FIELD, bordercolor=LINE,
                    lightcolor=ACCENT, darkcolor=ACCENT)
        s.configure("Vertical.TScrollbar", background=PANEL, troughcolor=FIELD, bordercolor=LINE,
                    arrowcolor=TEXT_DIM)

    def _init_vars(self):
        self.event_name_var = tk.StringVar(value="Oppvisningsstevne")
        self.location_var = tk.StringVar(value="Iskanten")
        self.start_date_var = tk.StringVar(value=datetime.now().strftime("%d.%m.%y"))
        self.start_time_var = tk.StringVar(value="18:00")
        self.interval_var = tk.StringVar(value="3:10")
        self.ready_var = tk.StringVar(value="0:30")
        self.group_count_var = tk.StringVar(value="2")
        self.group_size_var = tk.StringVar(value="8")
        self.warmup_var = tk.StringVar(value="4:00")
        self.water_count_var = tk.StringVar(value="0")
        self.water_duration_var = tk.StringVar(value="15:00")
        # Feltene som lagres i oppsettfilen.
        self.setup_vars = {
            "stevne": self.event_name_var,
            "sted": self.location_var,
            "dato": self.start_date_var,
            "start": self.start_time_var,
            "lengde_lop": self.interval_var,
            "klart_til_neste": self.ready_var,
            "antall_oppvarmingsgrupper": self.group_count_var,
            "utovere_per_gruppe": self.group_size_var,
            "oppvarming": self.warmup_var,
            "antall_vanningspauser": self.water_count_var,
            "lengde_vanningspause": self.water_duration_var,
        }
        for var in (self.event_name_var, self.location_var, self.start_date_var):
            var.trace_add("write", lambda *_: self.mark_dirty())
        for var in (
            self.start_time_var,
            self.interval_var,
            self.ready_var,
            self.warmup_var,
            self.water_duration_var,
        ):
            var.trace_add("write", lambda *_: self.refresh_schedule())
        self.group_count_var.trace_add("write", lambda *_: self.on_group_count_change())
        self.group_size_var.trace_add("write", lambda *_: self.on_layout_setting_change(regroup=True))
        self.water_count_var.trace_add("write", lambda *_: self.on_layout_setting_change(regroup=False))
        self.new_name_var = tk.StringVar()
        self.new_club_var = tk.StringVar()
        self.var_pdf = tk.BooleanVar(value=False)
        self.var_excel = tk.BooleanVar(value=True)
        self.var_html = tk.BooleanVar(value=True)

    def _build_menu(self):
        # Egen menylinje, siden Windows' innebygde menylinje ikke kan farges.
        menubar = tk.Frame(self.root, bg=BG)
        menubar.pack(side="top", fill="x", padx=(8, 0))
        tk.Frame(self.root, bg=LINE, height=1).pack(side="top", fill="x")
        self.menu_options = dict(
            tearoff=0, bg=PANEL, fg=TEXT, activebackground=ACCENT, activeforeground=ACCENT_TEXT,
            selectcolor=TEXT, disabledforeground=TEXT_DISABLED, relief="flat", bd=0,
            font=self.fonts["body"],
        )

        def add_menu(label):
            button = tk.Menubutton(menubar, text=label, underline=0, bg=BG, fg=TEXT,
                                   activebackground=BUTTON_HOVER, activeforeground=TEXT, relief="flat",
                                   bd=0, padx=10, pady=4, font=self.fonts["body"])
            button.pack(side="left")
            menu = tk.Menu(button, **self.menu_options)
            button.configure(menu=menu)
            return menu

        menu_file = add_menu("Fil")
        menu_file.add_command(label="Hent deltakerliste fra isonen…", command=self.load_isonen)
        menu_file.add_command(label="Hent deltakere fra musikktjenesten…", command=self.load_music_entries)
        menu_file.add_command(label="Velg musikk-zip…", command=self.choose_music_zip)
        menu_file.add_command(
            label="Skann mappe (isonen, musikktjenesten og musikk-zip)…", command=self.scan_folder
        )
        menu_file.add_separator()
        menu_file.add_command(label="Nytt oppsett", command=self.clear_rows)
        menu_file.add_command(label="Åpne oppsett…", accelerator="Ctrl+O", command=self.open_setup)
        menu_file.add_command(label="Lagre oppsett", accelerator="Ctrl+S", command=self.save_setup)
        menu_file.add_command(label="Lagre oppsett som…", command=lambda: self.save_setup(ask=True))
        menu_file.add_separator()
        menu_file.add_command(label="Avslutt", command=self.on_close)
        menu_reports = add_menu("Rapporter")
        menu_reports.add_command(label="Lagre startliste (Excel + PDF)", command=self.generate_startliste)
        menu_reports.add_separator()
        menu_reports.add_command(label="Lag deltakerrapport", command=self.generate_files)
        menu_reports.add_checkbutton(label="PDF", variable=self.var_pdf)
        menu_reports.add_checkbutton(label="Excel", variable=self.var_excel)
        menu_reports.add_checkbutton(label="HTML", variable=self.var_html)
        menu_settings = add_menu("Innstillinger")
        self.menu_tabellvisning = tk.Menu(menu_settings, **self.menu_options)
        menu_settings.add_cascade(label="Kolonner i stevnetabellen", menu=self.menu_tabellvisning)
        menu_help = add_menu("Hjelp")
        menu_help.add_command(label="Om", command=self.show_about)
        self.root.bind_all("<Control-s>", lambda e: self.save_setup())
        self.root.bind_all("<Control-o>", lambda e: self.open_setup())

    def _build(self):
        r = self.root
        r.title(APP_NAME)
        r.minsize(1100, 700)

        top = ttk.Frame(r, padding=(18, 12, 18, 6))
        top.pack(fill="x")
        title_box = ttk.Frame(top)
        title_box.pack(side="left", anchor="w")
        ttk.Label(title_box, text=APP_NAME, style="Title.TLabel").pack(anchor="w")
        self.status_var = tk.StringVar(
            value="Hent deltakere fra Fil-menyen, eller legg inn utøvere manuelt."
        )
        ttk.Label(title_box, textvariable=self.status_var, style="Dim.TLabel").pack(anchor="w")

        log_panel = ttk.Frame(r, style="Panel.TFrame", padding=(14, 8))
        log_panel.pack(side="bottom", fill="x", padx=18, pady=(0, 18))
        ttk.Label(log_panel, text="Logg", style="PanelHead.TLabel").pack(anchor="w")
        log_frame = ttk.Frame(log_panel, style="Panel.TFrame")
        log_frame.pack(fill="x", pady=(4, 0))
        self.log_widget = tk.Text(log_frame, height=5, bg=FIELD, fg=TEXT_DIM, insertbackground=TEXT,
                                  font=self.fonts["mono"], relief="flat", wrap="word", padx=10,
                                  pady=6, highlightthickness=0)
        log_scroll = ttk.Scrollbar(log_frame, command=self.log_widget.yview)
        self.log_widget.configure(yscrollcommand=log_scroll.set)
        log_scroll.pack(side="right", fill="y")
        self.log_widget.pack(side="left", fill="both", expand=True)

        body = ttk.Frame(r, padding=(18, 4, 18, 10))
        body.pack(fill="both", expand=True)
        nb = ttk.Notebook(body)
        nb.pack(fill="both", expand=True)
        self._tab_schedule(nb)
        self._tab_music(nb)

    def _panel(self, nb, title):
        frame = ttk.Frame(nb, style="Panel.TFrame", padding=20)
        nb.add(frame, text=title)
        return frame

    # -------------------- Fane: Tidsplan --------------------
    def _tab_schedule(self, nb):
        p = self._panel(nb, "Tidsplan/planlegging")
        p.columnconfigure(1, weight=1)
        p.rowconfigure(0, weight=1)

        form = ttk.Frame(p, style="Panel.TFrame")
        form.grid(row=0, column=0, sticky="ns", padx=(0, 24))
        ttk.Label(form, text="Oppsett", style="PanelHead.TLabel").grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        fields = [
            ("Stevne", self.event_name_var, 22, None),
            ("Sted", self.location_var, 22, None),
            ("Dato", self.start_date_var, 10, "date"),
            ("Start kl.", self.start_time_var, 8, None),
            ("Lengde på løp", self.interval_var, 8, None),
            ("Klart til neste løp", self.ready_var, 8, None),
            ("Antall oppvarmingsgrupper", self.group_count_var, 6, (0, 40)),
            ("Utøvere per gruppe", self.group_size_var, 6, (1, 40)),
            ("Oppvarming per gruppe", self.warmup_var, 8, None),
            ("Antall vanningspauser", self.water_count_var, 6, (0, 20)),
            ("Lengde på vanningspause", self.water_duration_var, 8, None),
        ]
        for i, (label, var, width, kind) in enumerate(fields, start=1):
            ttk.Label(form, text=label, style="Panel.TLabel").grid(
                row=i, column=0, sticky="w", pady=3, padx=(0, 12))
            widget = None
            if kind == "date":
                try:
                    from tkcalendar import DateEntry

                    widget = DateEntry(form, textvariable=var, date_pattern="dd.mm.yy", width=width)
                except Exception:
                    widget = None
            elif kind:
                widget = ttk.Spinbox(form, from_=kind[0], to=kind[1], textvariable=var, width=width,
                                     font=self.fonts["body"])
            if widget is None:
                widget = ttk.Entry(form, textvariable=var, width=width, font=self.fonts["body"])
            widget.grid(row=i, column=1, sticky="w", pady=3)
            if var is self.group_size_var:
                self.group_size_widget = widget
        self.update_group_size_state()
        row = len(fields) + 1
        ttk.Label(
            form,
            text="Tider skrives som M:SS, starttid som TT:MM.\n"
                 "Lengde på løp + klart til neste løp = tid per utøver.",
            style="PanelDim.TLabel",
            justify="left",
        ).grid(row=row, column=0, columnspan=2, sticky="w", pady=(4, 0))

        ttk.Label(form, text="Legg til utøver", style="PanelHead.TLabel").grid(
            row=row + 1, column=0, columnspan=2, sticky="w", pady=(22, 8))
        ttk.Label(form, text="Navn", style="Panel.TLabel").grid(row=row + 2, column=0, sticky="w", pady=3)
        name_entry = ttk.Entry(form, textvariable=self.new_name_var, width=22, font=self.fonts["body"])
        name_entry.grid(row=row + 2, column=1, sticky="w", pady=3)
        ttk.Label(form, text="Klubb", style="Panel.TLabel").grid(row=row + 3, column=0, sticky="w", pady=3)
        club_entry = ttk.Entry(form, textvariable=self.new_club_var, width=22, font=self.fonts["body"])
        club_entry.grid(row=row + 3, column=1, sticky="w", pady=3)
        for entry in (name_entry, club_entry):
            entry.bind("<Return>", lambda e: self.add_runner())
        self.name_entry = name_entry
        ttk.Button(form, text="Legg til", command=self.add_runner).grid(
            row=row + 4, column=1, sticky="w", pady=(6, 0))

        actions = ttk.Frame(form, style="Panel.TFrame")
        actions.grid(row=row + 5, column=0, columnspan=2, sticky="we", pady=(24, 0))
        # Import, åpne og lagre ligger i Fil-menyen, så knappene ikke klippes på små skjermer.
        ttk.Button(actions, text="Lagre startliste (Excel + PDF)", style="Accent.TButton",
                   command=self.generate_startliste).pack(fill="x")

        table = ttk.Frame(p, style="Panel.TFrame")
        table.grid(row=0, column=1, sticky="nsew")
        table.rowconfigure(0, weight=1)
        table.columnconfigure(0, weight=1)
        columns = {
            "nr": ("Nr.", 50, "center"),
            "start": ("Start", 90, "center"),
            "dash": ("", 24, "center"),
            "slutt": ("Slutt", 90, "center"),
            "lengde": ("Lengde på løp", 110, "center"),
            "tid": ("Sammenlagt tidsforbruk", 170, "center"),
            "navn": ("Navn", 260, "w"),
            "klubb": ("Klubb", 200, "w"),
        }
        self.schedule_tree = ttk.Treeview(table, columns=tuple(columns), show="headings",
                                          selectmode="browse")
        for col, (text, width, anchor) in columns.items():
            self.schedule_tree.heading(col, text=text, anchor=anchor)
            self.schedule_tree.column(col, width=width, anchor=anchor, stretch=col in ("navn", "klubb"))
        self.schedule_tree.tag_configure("warmup", background=GROUP_BG, font=self.fonts["bold"])
        self.schedule_tree.tag_configure("pause", background=GROUP_BG, foreground=WARN,
                                         font=self.fonts["bold"])
        self.schedule_tree.tag_configure("overfull", background=WARN, foreground=ACCENT_TEXT,
                                         font=self.fonts["bold"])
        self.schedule_tree.bind("<Motion>", self.on_schedule_hover)
        self.schedule_tree.bind("<Leave>", lambda e: self.hide_tip())
        self.schedule_tree.bind("<Double-1>", lambda e: self.edit_selected())
        self.schedule_tree.bind("<Delete>", lambda e: self.delete_selected())
        self.schedule_tree.bind("<Button-3>", self.on_schedule_right_click)
        self.schedule_menu = tk.Menu(self.schedule_tree, **self.menu_options)
        self.schedule_menu.add_command(label="Endre lengde på løp…", command=self.edit_run_length)
        self.schedule_menu.add_command(label="Bruk standard lengde", command=self.reset_run_length)
        self.schedule_menu.add_separator()
        self.schedule_menu.add_command(label="Rediger", command=self.edit_selected)
        self.schedule_menu.add_command(label="Slett", command=self.delete_selected)
        self.schedule_tree.bind("<Alt-Up>", lambda e: self.move_selected(-1) or "break")
        self.schedule_tree.bind("<Alt-Down>", lambda e: self.move_selected(1) or "break")
        self.schedule_tree.bind("<ButtonPress-1>", self.on_drag_start)
        self.schedule_tree.bind("<B1-Motion>", self.on_drag_motion)
        self.schedule_tree.bind("<ButtonRelease-1>", self.on_drag_drop)
        self.drop_line = tk.Frame(self.schedule_tree, bg=ACCENT, height=3)
        scroll = ttk.Scrollbar(table, orient="vertical", command=self.schedule_tree.yview)
        self.schedule_tree.configure(yscrollcommand=scroll.set)
        self.schedule_tree.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")

        self.summary_var = tk.StringVar()
        self.summary_label = ttk.Label(table, textvariable=self.summary_var, style="PanelDim.TLabel")
        self.summary_label.grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))

        side = ttk.Frame(table, style="Panel.TFrame")
        side.grid(row=0, column=2, sticky="ns", padx=(12, 0))
        for text, command, gap in (
            ("Flytt opp", lambda: self.move_selected(-1), 0),
            ("Flytt ned", lambda: self.move_selected(1), 4),
            ("Rediger", self.edit_selected, 16),
            ("Slett", self.delete_selected, 4),
            ("Randomiser", self.shuffle_rows, 16),
            ("Sorter fornavn", self.sort_by_given, 4),
            ("Sorter etternavn", self.sort_by_family, 4),
            ("Eksporter til Excel", self.export_excel, 16),
        ):
            ttk.Button(side, text=text, command=command, width=16).pack(fill="x", pady=(gap, 0))
        ttk.Label(side, text="Dra rader med musen,\neller bruk Alt+↑/↓.\nDobbeltklikk redigerer.\n"
                             "Høyreklikk en utøver for\nå endre lengde på løp.",
                  style="PanelDim.TLabel", justify="left").pack(anchor="w", pady=(16, 0))

    # -------------------- Fane: Stevne --------------------
    def _tab_music(self, nb):
        p = self._panel(nb, "Stevne")
        p.columnconfigure(0, weight=1)
        p.rowconfigure(2, weight=1)

        # Bahnschrift har ulik bredde på sifrene, så hvert tegn får en fast plass for å unngå at klokka hopper.
        clock_font = tkfont.Font(font=self.fonts["clock"])
        digit_w = max(clock_font.measure(d) for d in "0123456789")
        colon_w = clock_font.measure(":")
        widths = [digit_w, digit_w, colon_w, digit_w, digit_w, colon_w, digit_w, digit_w]
        self.clock_canvas = tk.Canvas(p, width=sum(widths), height=clock_font.metrics("linespace"),
                                      bg=PANEL, highlightthickness=0)
        self.clock_canvas.grid(row=0, column=0, pady=(0, 12))
        self.clock_items = []
        x = 0
        for width in widths:
            self.clock_items.append(self.clock_canvas.create_text(
                x + width / 2, 0, anchor="n", text="0", fill=TEXT, font=clock_font))
            x += width

        bar = ttk.Frame(p, style="Panel.TFrame")
        bar.grid(row=1, column=0, sticky="we", pady=(0, 12))
        ttk.Button(bar, text="Velg musikk-zip…", style="Accent.TButton",
                   command=self.choose_music_zip).pack(side="left")
        ttk.Button(bar, text="Lag spilleliste", command=self.generate_playlist_only).pack(
            side="left", padx=(8, 0))
        ttk.Checkbutton(bar, text="Bruk VLC direkte", variable=self.use_external_player_var,
                        command=self.on_use_external_player_toggle).pack(side="left", padx=(16, 0))
        self.music_info_var = tk.StringVar(value="Ingen musikk-zip valgt.")
        ttk.Label(bar, textvariable=self.music_info_var, style="PanelDim.TLabel").pack(
            side="left", padx=(16, 0))

        table = ttk.Frame(p, style="Panel.TFrame")
        table.grid(row=2, column=0, sticky="nsew")
        table.rowconfigure(0, weight=1)
        table.columnconfigure(0, weight=1)
        columns = {
            "startnummer": "Nr.",
            "lop_gjennomfort": "Løp gjennomført",
            "starttid": "Start",
            "navn": "Navn",
            "klubb": "Klubb",
            "klasse": "Klasse",
            "program": "Program",
            "påmelding": "Påmelding",
            "musikkstatus": "Musikkstatus",
            "musikknavn": "Lydfil",
            "musikkformat": "Format",
            "musikktid": "Lengde",
        }
        self.tree = ttk.Treeview(table, columns=tuple(columns), show="headings", selectmode="browse")
        for col, text in columns.items():
            left = ("navn", "klubb", "klasse", "påmelding", "musikkstatus", "musikknavn")
            anchor = "w" if col in left else "center"
            self.tree.heading(col, text=text, anchor=anchor)
            self.tree.column(col, width=100, anchor=anchor)
        self.tree.tag_configure("missing_music", foreground=BAD)
        self.tree.tag_configure("audio_warning", foreground=WARN)
        self.tree.tag_configure("audio_error", foreground=BAD, background="#3B1D24")
        self.tree.tag_configure("cancelled", foreground=TEXT_DISABLED)
        self.tree.tag_configure("completed", foreground=GOOD, background="#16302A")
        self.tree.bind("<Button-1>", self.on_tree_checkbox_click)
        self.tree.bind("<Double-1>", self.on_tree_double_click)
        self.tree.bind("<Return>", lambda e: self.toggle_play_pause())
        scroll = ttk.Scrollbar(table, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")
        ttk.Label(p, text="Dobbeltklikk eller Enter spiller lydfilen. Klikk i «Løp gjennomført» for å krysse av.",
                  style="PanelDim.TLabel").grid(row=3, column=0, sticky="w", pady=(8, 0))

        self.table_columns = tuple(columns)
        self.table_column_vars = {}
        for column, text in columns.items():
            variable = tk.BooleanVar(value=True)
            self.table_column_vars[column] = variable
            self.menu_tabellvisning.add_checkbutton(
                label=text,
                variable=variable,
                command=lambda changed=column: self.update_visible_table_columns(changed),
            )
        self.menu_tabellvisning.add_separator()
        self.menu_tabellvisning.add_command(label="Vis alle kolonner", command=self.show_all_table_columns)

        # Avspiller i kort med runde knapper og lydbølge man kan klikke i for å spole.
        card = tk.Frame(p, bg=FIELD, highlightbackground=LINE, highlightthickness=1, padx=14, pady=10)
        card.grid(row=4, column=0, sticky="we", pady=(16, 0))
        player = tk.Frame(card, bg=FIELD)
        player.pack(fill="x")
        self.player_buttons = {}
        for name, size, command in (
            ("back", 34, lambda: self.skip_seconds(-SKIP_SECONDS)),
            ("play", 50, self.toggle_play_pause),
            ("forward", 34, lambda: self.skip_seconds(SKIP_SECONDS)),
            ("stop", 34, self.stop_playback),
        ):
            canvas = tk.Canvas(player, width=size, height=size, bg=FIELD, highlightthickness=0,
                               cursor="hand2")
            canvas.pack(side="left", padx=(0, 8))
            canvas.bind("<Button-1>", lambda e, c=command: self.player_enabled and c())
            canvas.bind("<Enter>", lambda e, n=name: self.draw_player_button(n, hover=True))
            canvas.bind("<Leave>", lambda e, n=name: self.draw_player_button(n))
            self.player_buttons[name] = canvas
        self.player_time_var = tk.StringVar(value="0:00")
        self.player_total_var = tk.StringVar(value="0:00")
        ttk.Label(player, textvariable=self.player_total_var, background=FIELD,
                  font=self.fonts["head"], foreground=ACCENT).pack(side="right")
        ttk.Label(player, text=" / ", background=FIELD, foreground=TEXT_DIM).pack(side="right")
        ttk.Label(player, textvariable=self.player_time_var, background=FIELD,
                  foreground=TEXT_DIM).pack(side="right")
        self.wave_canvas = tk.Canvas(player, height=50, bg=FIELD, highlightthickness=0, cursor="hand2")
        self.wave_canvas.pack(side="left", fill="x", expand=True, padx=(6, 14))
        self.wave_canvas.bind("<Configure>", lambda e: self.draw_waveform())
        self.wave_canvas.bind("<Button-1>", self.on_waveform_click)
        self.wave_canvas.bind("<B1-Motion>", self.on_waveform_click)
        self.player_track_var = tk.StringVar(value="Ingen spor valgt")
        ttk.Label(card, textvariable=self.player_track_var, background=FIELD,
                  foreground=TEXT_DIM, font=self.fonts["small"]).pack(anchor="w", pady=(8, 0))
        self.play_pause_text.trace_add("write", lambda *_: self.draw_player_button("play"))
        for name in self.player_buttons:
            self.draw_player_button(name)

    # -------------------- Generelt --------------------
    def log(self, msg):
        self.log_widget.insert("end", msg + "\n")
        self.log_widget.see("end")
        self.root.update_idletasks()

    def update_clock(self):
        for item, char in zip(self.clock_items, datetime.now().strftime("%H:%M:%S")):
            self.clock_canvas.itemconfigure(item, text=char)
        self.root.after(500, self.update_clock)

    def output_dir(self):
        out_dir = (self.source_dir or Path(__file__).resolve().parent) / "output"
        out_dir.mkdir(parents=True, exist_ok=True)
        return out_dir

    def active_rows(self):
        return [r for r in self.rows if not is_cancelled(r.get("Påmelding", ""))]

    def visible_items(self):
        return [i for i in self.items if is_marker(i) or not is_cancelled(i.get("Påmelding", ""))]

    def index_of(self, item):
        return next((i for i, x in enumerate(self.items) if x is item), None)

    def confirm_replace(self):
        if self.dirty:
            return self.confirm_discard()
        if not self.items:
            return True
        return messagebox.askyesno(
            "Erstatte listen?",
            "Dette erstatter den nåværende listen. Vil du fortsette?",
            parent=self.root,
        )

    def confirm_discard(self):
        if not self.dirty or not self.items:
            return True
        answer = messagebox.askyesnocancel(
            "Ulagrede endringer",
            "Oppsettet har endringer som ikke er lagret. Vil du lagre først?",
            parent=self.root,
        )
        if answer is None:
            return False
        return self.save_setup() if answer else True

    def mark_dirty(self):
        if self._loading:
            return
        self.dirty = True
        self.update_title()

    def update_title(self):
        name = self.setup_path.name if self.setup_path else "Nytt oppsett"
        self.root.title(f"{APP_NAME} – {name}{' *' if self.dirty else ''}")

    def update_status(self):
        parts = []
        if self.setup_path:
            parts.append(f"Oppsett: {self.setup_path}")
        elif self.source_dir:
            parts.append(f"Kilde: {self.source_dir}")
        parts.append(f"{len(self.active_rows())} utøvere")
        if self.music_zip:
            parts.append(f"Musikk: {self.music_zip.name}")
        self.status_var.set(" · ".join(parts))

    # -------------------- Tidsplan --------------------
    @staticmethod
    def parse_optional_duration(value):
        value = (value or "").strip()
        if value in ("", "0"):
            return 0
        return parse_duration_mmss(value)

    @staticmethod
    def read_int(var, default):
        try:
            return int(var.get())
        except (ValueError, tk.TclError):
            return default

    def read_schedule_settings(self):
        start = parse_time_hhmm(self.start_time_var.get())
        if not start:
            return "Ugyldig starttid. Bruk TT:MM."
        run_seconds = parse_duration_mmss(self.interval_var.get())
        if not run_seconds:
            return "Ugyldig lengde på løp. Bruk M:SS."
        ready = self.parse_optional_duration(self.ready_var.get())
        if ready is None:
            return "Ugyldig «klart til neste løp». Bruk M:SS."
        warmup = self.parse_optional_duration(self.warmup_var.get())
        if warmup is None:
            return "Ugyldig oppvarming. Bruk M:SS."
        water_seconds = self.parse_optional_duration(self.water_duration_var.get())
        if water_seconds is None:
            return "Ugyldig lengde på vanningspause. Bruk M:SS."
        return {
            "run_seconds": run_seconds,
            "start_time": start,
            "ready_seconds": ready,
            "warmup_seconds": warmup,
            "water_seconds": water_seconds,
        }

    def group_count(self):
        return max(0, self.read_int(self.group_count_var, 0))

    def group_limit(self):
        # Maks utøvere per gruppe. Med 0 eller 1 gruppe er det ingen grense.
        return self.read_int(self.group_size_var, 8) if self.group_count() > 1 else float("inf")

    def set_auto_group_size(self):
        self._loading = True
        try:
            self.group_size_var.set(str(auto_group_size(len(self.active_rows()), self.group_count())))
        finally:
            self._loading = False

    def update_group_size_state(self):
        # «Utøvere per gruppe» kan bare velges når det er minst 2 grupper.
        if hasattr(self, "group_size_widget"):
            self.group_size_widget.configure(state="normal" if self.group_count() > 1 else "disabled")

    def relayout(self, auto_size=False):
        if auto_size:
            self.set_auto_group_size()
        self.items = layout_items(
            self.rows,
            self.group_count(),
            self.read_int(self.group_size_var, 8),
            self.read_int(self.water_count_var, 0),
        )

    def on_group_count_change(self):
        self.update_group_size_state()
        if self._loading:
            return
        self.relayout(auto_size=True)
        self.refresh_schedule()

    def on_layout_setting_change(self, regroup):
        if self._loading:
            return
        if regroup:
            self.relayout()
        else:
            self.items = place_water_pauses(self.items, self.read_int(self.water_count_var, 0))
        self.refresh_schedule()

    def sync_marker_counts(self):
        # Oppdaterer antall grupper og vanningspauser etter det som faktisk står i tidsplanen.
        self._loading = True
        try:
            self.group_count_var.set(str(sum(1 for i in self.items if i.get("Kind") == "warmup")))
            self.water_count_var.set(str(sum(1 for i in self.items if i.get("Kind") == "pause")))
        finally:
            self._loading = False

    def selected_item(self):
        selected = self.schedule_tree.selection()
        if not selected:
            return None
        # Valget kan være igjen fra før listen ble byttet ut eller kortet ned.
        index = int(selected[0][1:])
        return self.items[index] if index < len(self.items) else None

    def refresh_schedule(self, select=None):
        if not hasattr(self, "schedule_tree"):
            return
        if select is None:
            select = self.selected_item()
        self.mark_dirty()
        settings = self.read_schedule_settings()
        if isinstance(settings, str):
            self.summary_label.configure(style="PanelWarn.TLabel")
            self.summary_var.set(settings)
            return
        self.summary_label.configure(style="PanelDim.TLabel")

        self.entries = build_startliste(self.items, **settings)
        for row in self.rows:
            row["StartTid"] = ""
            row["SluttTid"] = ""
        tree = self.schedule_tree
        tree.delete(*tree.get_children())
        positions = {id(item): i for i, item in enumerate(self.items)}
        group_size = self.group_limit()
        overflows = group_overflows(self.items, group_size)
        self.overflow_tips = {}
        select_iid = None
        for entry in self.entries:
            item = entry["row"]
            iid = f"i{positions[id(item)]}"
            if item is select:
                select_iid = iid
            if entry["kind"] == "run":
                item["StartTid"] = entry["start"]
                item["SluttTid"] = entry["end"]
            tags = () if entry["kind"] == "run" else (entry["kind"],)
            count = overflows.get(id(item))
            if count:
                tags = ("overfull",)
                extra = count - group_size
                self.overflow_tips[iid] = (
                    f"{entry['navn']} har {count} utøvere, men maks er {group_size}.\n"
                    f"Flytt {extra} utøver{'e' if extra > 1 else ''} til en annen gruppe, "
                    "eller øk «Utøvere per gruppe»."
                )
            run_length = entry["run_length"] + (" *" if entry["custom_length"] else "")
            tree.insert("", "end", iid=iid, tags=tags, values=(
                entry["nr"], entry["start"], "-", entry["end"], run_length, entry["duration"],
                entry["navn"], entry["klubb"],
            ))
        if select_iid:
            tree.selection_set(select_iid)
            tree.focus(select_iid)
            tree.see(select_iid)

        runners = sum(1 for e in self.entries if e["kind"] == "run")
        if not runners:
            self.summary_var.set("Ingen utøvere ennå. Hent deltakere fra Fil-menyen eller legg til manuelt.")
        else:
            groups = sum(1 for e in self.entries if e["kind"] == "warmup")
            pauses = sum(1 for e in self.entries if e["kind"] == "pause")
            parts = [f"{runners} utøvere"]
            if groups:
                parts.append(f"{groups} oppvarmingsgruppe{'r' if groups > 1 else ''}")
            parts += [f"{pauses} vanningspauser", f"ferdig kl. {self.entries[-1]['end']}"]
            text = " · ".join(parts)
            wanted = self.read_int(self.water_count_var, 0)
            if pauses < wanted:
                where = "mellom gruppene" if groups else "mellom utøverne"
                text += f" (bare {pauses} vanningspauser får plass {where})"
            if overflows:
                text += f" · {len(overflows)} gruppe{'r' if len(overflows) > 1 else ''} har for mange utøvere"
                self.summary_label.configure(style="PanelWarn.TLabel")
            self.summary_var.set(text)
        self.update_status()
        self.refresh_table()

    def add_runner(self):
        given, family = split_name(self.new_name_var.get().strip())
        if not given:
            self.name_entry.focus_set()
            return
        row = new_participant(given, family, self.new_club_var.get())
        # Utøveren havner i siste gruppe. Finnes ingen grupper ennå, lages de på nytt.
        self.items.append(row)
        if self.group_count() and not any(i.get("Kind") == "warmup" for i in self.items):
            self.relayout(auto_size=True)
        self.match_missing_music()
        self.new_name_var.set("")
        self.name_entry.focus_set()
        self.log(f"La til utøver: {display_name(row)}")
        self.refresh_schedule(select=row)

    def edit_selected(self):
        row = self.selected_item()
        if row is None or is_marker(row):
            return
        name = simpledialog.askstring("Rediger utøver", "Navn:", initialvalue=display_name(row),
                                      parent=self.root)
        if name is None or not name.strip():
            return
        club = simpledialog.askstring("Rediger utøver", "Klubb:",
                                      initialvalue=row.get("Organisation", ""), parent=self.root)
        if club is None:
            return
        given, family = split_name(name.strip())
        if (given, family) != (row.get("GivenName"), row.get("FamilyName")):
            row["MusikkFil"] = ""
        row["GivenName"], row["FamilyName"] = given, family
        row["NavnFraIsonen"] = f"{given} {family}".strip()
        row["PrintName"] = f"{family}, {given}".strip(", ")
        row["Organisation"] = club.strip()
        self.match_missing_music()
        self.refresh_schedule(select=row)

    def on_schedule_right_click(self, event):
        iid = self.schedule_tree.identify_row(event.y)
        if not iid:
            return
        self.schedule_tree.selection_set(iid)
        self.schedule_tree.focus(iid)
        row = self.selected_item()
        if row is None or is_marker(row):
            return
        self.hide_tip()
        self.schedule_menu.entryconfigure(1, state="normal" if row.get("LopLengde") else "disabled")
        try:
            self.schedule_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.schedule_menu.grab_release()

    def edit_run_length(self):
        row = self.selected_item()
        if row is None or is_marker(row):
            return
        standard = parse_duration_mmss(self.interval_var.get())
        value = simpledialog.askstring(
            "Lengde på løp",
            f"Lengde på løp for {display_name(row)} (M:SS):",
            initialvalue=format_duration(row.get("LopLengde") or standard),
            parent=self.root,
        )
        if value is None:
            return
        seconds = parse_duration_mmss(value.strip())
        if not seconds:
            messagebox.showerror("Ugyldig lengde", "Ugyldig lengde på løp. Bruk M:SS.", parent=self.root)
            return
        if seconds == standard:
            row.pop("LopLengde", None)
        else:
            row["LopLengde"] = seconds
        self.log(f"Lengde på løp for {display_name(row)}: {format_duration(seconds)}")
        self.refresh_schedule(select=row)

    def reset_run_length(self):
        row = self.selected_item()
        if row is None or is_marker(row) or not row.pop("LopLengde", None):
            return
        self.log(f"{display_name(row)} bruker standard lengde på løp.")
        self.refresh_schedule(select=row)

    def move_item(self, src, dst):
        # Etter pop havner raden under målet når den flyttes ned, og over når den flyttes opp.
        group_size = self.group_limit()
        before = group_overflows(self.items, group_size)
        item = self.items.pop(src)
        self.items.insert(dst, item)
        self.refresh_schedule(select=item)
        after = group_overflows(self.items, group_size)
        grown = [key for key, count in after.items() if count > before.get(key, 0)]
        if grown:
            iids = {f"i{k}": id(x) for k, x in enumerate(self.items)}
            tips = [tip for iid, tip in self.overflow_tips.items() if iids.get(iid) in grown]
            messagebox.showwarning("For mange utøvere i gruppen", "\n\n".join(tips), parent=self.root)

    def on_schedule_hover(self, event):
        tip = self.overflow_tips.get(self.schedule_tree.identify_row(event.y))
        if not tip:
            self.hide_tip()
            return
        if self.tip is None:
            self.tip = tk.Toplevel(self.root)
            self.tip.overrideredirect(True)
            self.tip.attributes("-topmost", True)
            self.tip_label = tk.Label(self.tip, bg=PANEL, fg=TEXT, font=self.fonts["body"], justify="left",
                                      wraplength=380, padx=10, pady=6, highlightthickness=2,
                                      highlightbackground=WARN)
            self.tip_label.pack()
        self.tip_label.configure(text=tip)
        self.tip.geometry(f"+{event.x_root + 14}+{event.y_root + 18}")
        self.tip.deiconify()

    def hide_tip(self):
        if self.tip is not None:
            self.tip.withdraw()

    def move_selected(self, delta):
        item = self.selected_item()
        if item is None:
            messagebox.showinfo("Info", "Velg en rad i tabellen først.")
            return
        visible = self.visible_items()
        pos = next(i for i, x in enumerate(visible) if x is item)
        target = pos + delta
        if 0 <= target < len(visible):
            self.move_item(self.index_of(item), self.index_of(visible[target]))

    def on_drag_start(self, event):
        self.hide_tip()
        self.drag_iid = self.schedule_tree.identify_row(event.y) or None

    def on_drag_motion(self, event):
        tree = self.schedule_tree
        if not self.drag_iid:
            return
        if event.y < 30:
            tree.yview_scroll(-1, "units")
        elif event.y > tree.winfo_height() - 10:
            tree.yview_scroll(1, "units")
        target = tree.identify_row(event.y)
        bbox = tree.bbox(target) if target and target != self.drag_iid else None
        if not bbox:
            self.drop_line.place_forget()
            return
        _, y, _, height = bbox
        below = tree.index(target) > tree.index(self.drag_iid)
        self.drop_line.place(x=0, y=y + height - 2 if below else y - 1, relwidth=1)
        tree.configure(cursor="sb_v_double_arrow")

    def on_drag_drop(self, event):
        self.drop_line.place_forget()
        self.schedule_tree.configure(cursor="")
        source, self.drag_iid = self.drag_iid, None
        target = self.schedule_tree.identify_row(event.y)
        if source and target and source != target:
            self.move_item(int(source[1:]), int(target[1:]))

    def delete_selected(self):
        item = self.selected_item()
        if item is None:
            messagebox.showinfo("Info", "Velg en rad i tabellen først.")
            return
        if item.get("Kind") == "pause":
            navn, question = "vanningspause", "Slette valgt vanningspause?"
        elif item.get("Kind") == "warmup":
            navn, question = "oppvarmingsgruppe", "Slette valgt oppvarmingsgruppe?\n\nUtøverne blir stående."
        else:
            navn = display_name(item)
            question = f"Slette valgt utøver?\n\n{navn}"
        if not messagebox.askyesno("Bekreft sletting", question):
            return
        visible = self.visible_items()
        pos = next(i for i, x in enumerate(visible) if x is item)
        del self.items[self.index_of(item)]
        self.sync_marker_counts()
        self.log(f"Slettet {navn}.")
        remaining = self.visible_items()
        self.refresh_schedule(select=remaining[min(pos, len(remaining) - 1)] if remaining else None)

    def set_runner_order(self, runners):
        # Bytter rekkefølgen på aktive utøvere; grupper, pauser og avmeldte står på sine plasser.
        slots = [k for k, item in enumerate(self.items)
                 if not is_marker(item) and not is_cancelled(item.get("Påmelding", ""))]
        for slot, row in zip(slots, runners):
            self.items[slot] = row
        self.refresh_schedule()

    def shuffle_rows(self):
        runners = self.active_rows()
        random.shuffle(runners)
        self.set_runner_order(runners)

    def sort_by_given(self):
        self.set_runner_order(sorted(
            self.active_rows(),
            key=lambda r: (normalize_text(r.get("GivenName")), normalize_text(r.get("FamilyName"))),
        ))

    def sort_by_family(self):
        self.set_runner_order(sorted(
            self.active_rows(),
            key=lambda r: (normalize_text(r.get("FamilyName")), normalize_text(r.get("GivenName"))),
        ))

    def clear_rows(self):
        if not self.confirm_discard():
            return
        self.items = []
        self.music_entries = []
        self.reset_music()
        self.setup_path = None
        self.log("Startet et nytt oppsett.")
        self.refresh_schedule()
        self.dirty = False
        self.update_title()

    # -------------------- Oppsettfil --------------------
    def save_setup(self, ask=False):
        path = self.setup_path
        if ask or not path:
            date_obj = parse_date_ddmmyy(self.start_date_var.get())
            date_part = date_obj.strftime("%Y-%m-%d") if date_obj else "dato"
            location = self.location_var.get().strip() or "sted"
            path = filedialog.asksaveasfilename(
                title="Lagre oppsett",
                defaultextension=".json",
                filetypes=SETUP_FILETYPES,
                initialfile=sanitize_filename(f"Oppsett_{location}_{date_part}") + ".json",
                initialdir=str(path.parent if path else self.output_dir()),
            )
            if not path:
                return False
            path = Path(path)
        data = {
            "type": "kunstlopassistenten-oppsett",
            "version": 1,
            "saved": datetime.now().isoformat(timespec="seconds"),
            "settings": {key: var.get() for key, var in self.setup_vars.items()},
            "source_dir": str(self.source_dir) if self.source_dir else "",
            "music_zip": str(self.music_zip) if self.music_zip else "",
            "items": self.items,
        }
        # Skriver til en midlertidig fil først, så forrige oppsett ikke ødelegges hvis skrivingen feiler.
        tmp_path = path.with_name(path.name + ".tmp")
        try:
            tmp_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(tmp_path, path)
        except Exception as exc:
            try:
                tmp_path.unlink(missing_ok=True)
            except OSError:
                pass
            messagebox.showerror("Feil", f"Kunne ikke lagre oppsettet:\n{exc}")
            return False
        self.setup_path = path
        self.dirty = False
        self.update_title()
        self.update_status()
        self.log(f"Lagret oppsett: {path}")
        return True

    def open_setup(self):
        if not self.confirm_discard():
            return
        path = filedialog.askopenfilename(
            title="Åpne oppsett",
            filetypes=SETUP_FILETYPES,
            initialdir=str(self.setup_path.parent if self.setup_path else self.output_dir()),
        )
        if not path:
            return
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8"))
        except Exception as exc:
            messagebox.showerror("Feil", f"Kunne ikke lese oppsettet:\n{exc}")
            return
        if not isinstance(data, dict) or not isinstance(data.get("items"), list):
            if isinstance(data, dict) and isinstance(data.get("order"), list):
                self.load_legacy_order(data["order"], path)
            else:
                messagebox.showerror("Feil", "Filen er ikke et oppsett fra Kunstløpassistenten.")
            return
        error = validate_setup(data)
        if error:
            messagebox.showerror("Feil", f"Kunne ikke åpne oppsettet:\n{error}")
            return

        self._loading = True
        try:
            for key, value in (data.get("settings") or {}).items():
                if key in self.setup_vars:
                    self.setup_vars[key].set(value)
        finally:
            self._loading = False
        self.items = list(data["items"])
        # Gamle oppsett mangler antall grupper; tidsplanen i filen bestemmer.
        self.sync_marker_counts()
        self.source_dir = Path(data["source_dir"]) if data.get("source_dir") else None
        self.music_entries = []
        # Musikk fra forrige oppsett skal aldri henge igjen.
        self.reset_music()
        music = data.get("music_zip")
        if music and Path(music).is_file():
            self.load_music_zip(music, match=False)
            self.match_missing_music()
        elif music:
            self.log(f"Fant ikke musikk-zip fra oppsettet: {music}")
        self.setup_path = Path(path)
        self.refresh_schedule()
        self.dirty = False
        self.update_title()
        self.log(f"Åpnet oppsett: {path}")

    # -------------------- Innlesing --------------------
    def load_isonen(self):
        initial = self.source_dir or (Path.home() / "Downloads")
        path = filedialog.askopenfilename(
            title="Velg deltakerliste fra isonen",
            initialdir=str(initial),
            filetypes=[("Excel", "*.xlsx"), ("Alle filer", "*.*")],
        )
        if not path or not self.confirm_replace():
            return
        rows = load_participants_from_excel(path, self.log)
        if not rows:
            messagebox.showerror("Feil", "Fant ingen deltakere i excel-filen.")
            return
        self.use_participants(rows, Path(path).parent)

    def load_music_entries(self):
        initial = self.source_dir or (Path.home() / "Downloads")
        path = filedialog.askopenfilename(
            title="Velg deltakerliste fra musikktjenesten (entries_….xlsx)",
            initialdir=str(initial),
            filetypes=[("Excel", "*.xlsx"), ("Alle filer", "*.*")],
        )
        if not path:
            return
        entries, info = load_entries_from_excel(path, self.log)
        if not entries:
            messagebox.showerror("Feil", "Fant ingen deltakere i excel-filen fra musikktjenesten.")
            return
        self.music_entries = entries
        self.apply_event_info(info)
        if self.rows:
            # Utfyller deltakerlisten som allerede er lastet (isonen eller oppsett).
            self.apply_music_entries()
            added = add_missing_segments(self.items)
            if added:
                self.log(f"{added} utøvere har også kortprogram. Friløpene er lagt bakerst i tidsplanen.")
            if self.music_files:
                self.match_all_music()
            self.refresh_schedule()
        elif self.confirm_replace():
            self.use_participants([participant_from_entry(e) for e in entries], Path(path).parent)

    def use_participants(self, rows, source_dir):
        # Ny deltakerliste: kobles mot musikktjenesten, deles i kortprogram/friløp og kobles mot musikk-zip.
        # Musikktjenesten flettes inn først, siden den kan ha øvelsen som mangler i isonen.
        self.rows = rows
        if self.music_entries:
            self.apply_music_entries()
        self.rows = split_segments(self.rows)
        self.setup_path = None
        self.source_dir = source_dir
        cancelled = len(self.rows) - len(self.active_rows())
        if cancelled:
            self.log(f"{cancelled} avmeldte er tatt ut av tidsplanen.")
        programs = sum(1 for r in self.rows if r.get("Program") == SEGMENT_NAMES["SP"])
        if programs:
            self.log(f"{programs} utøvere har kortprogram og friløp, og får ett løp for hvert program.")
        if self.music_files:
            self.match_all_music()
        self.relayout(auto_size=True)
        self.refresh_schedule()

    def apply_event_info(self, info):
        if info.get("title"):
            self.event_name_var.set(info["title"])
        if info.get("date"):
            self.start_date_var.set(info["date"].strftime("%d.%m.%y"))
        if info:
            date = f" {info['date']:%d.%m.%Y}" if info.get("date") else ""
            self.log(f"Stevne fra musikktjenesten: {info.get('title', '')}{date}")

    def apply_music_entries(self):
        # Kobler deltakerne fra musikktjenesten til radene med NIF ID (= Person ID i isonen), ellers navn.
        by_nif, by_name = {}, {}
        for row in self.rows:
            if row.get("NifId"):
                by_nif.setdefault(row["NifId"], []).append(row)
            by_name.setdefault(full_name_key(row), []).append(row)
        found = set()
        for entry in self.music_entries:
            rows = by_nif.get(entry["NifId"]) or by_name.get(full_name_key(entry), [])
            if not rows:
                self.log(f"I musikktjenesten, men ikke i deltakerlisten: "
                         f"{entry['GivenName']} {entry['FamilyName']} ({entry['Klasse']})")
                continue
            for row in rows:
                found.add(id(row))
                for field in ("NifId", "Klasse", "Ovelse"):
                    if not row.get(field):
                        row[field] = entry[field]
        missing = [r for r in self.active_rows() if id(r) not in found]
        for row in missing:
            self.log(f"Ikke påmeldt i musikktjenesten: {run_name(row)}")
        self.log(f"Musikktjenesten: {len(found)} løp koblet, {len(missing)} mangler.")

    def choose_music_zip(self):
        initial = self.source_dir or (Path.home() / "Downloads")
        path = filedialog.askopenfilename(
            title="Velg musikk-zip fra musikktjenesten",
            initialdir=str(initial),
            filetypes=[("Zip", "*.zip"), ("Alle filer", "*.*")],
        )
        if path and self.load_music_zip(path):
            self.refresh_schedule()

    def scan_folder(self):
        path = filedialog.askdirectory(
            title="Velg mappe med isonen-liste, deltakere fra musikktjenesten og musikk-zip",
            initialdir=str(self.source_dir or (Path.home() / "Downloads")),
        )
        if path and self.confirm_replace():
            self.read_folder(Path(path))

    @staticmethod
    def newest_file(folder, *patterns):
        files = {f for pattern in patterns for f in folder.glob(pattern) if f.is_file()}
        return max(files, key=lambda f: f.stat().st_mtime) if files else None

    def read_folder(self, folder):
        # Bruker nyeste fil av hver type, siden nedlastingsmappen ofte har flere versjoner.
        self.log_widget.delete("1.0", "end")
        isonen = self.newest_file(folder, "Deltakerliste*.xlsx")
        entries_path = self.newest_file(folder, "entries*.xlsx")
        music_zip = self.newest_file(folder, "Music*.zip", "Musikk*.zip")
        if not isonen and not entries_path:
            messagebox.showerror(
                "Feil",
                "Fant verken Deltakerliste*.xlsx (isonen) eller entries*.xlsx (musikktjenesten) i mappen.",
            )
            return
        self.music_entries = []
        if entries_path:
            self.music_entries, info = load_entries_from_excel(entries_path, self.log)
            self.apply_event_info(info)
        else:
            self.log("Fant ingen entries*.xlsx fra musikktjenesten.")
        if isonen:
            rows = load_participants_from_excel(isonen, self.log)
        else:
            self.log("Fant ingen Deltakerliste*.xlsx fra isonen. Bruker deltakerne fra musikktjenesten.")
            rows = [participant_from_entry(e) for e in self.music_entries]
        if not rows:
            messagebox.showerror("Feil", "Fant ingen deltakere i excel-filen.")
            return
        self.reset_music()
        if music_zip:
            self.load_music_zip(music_zip, match=False)
        else:
            self.log("Fant ingen Music*.zip eller Musikk*.zip i mappen.")
        self.use_participants(rows, folder)
        self.log(f"Totalt {len(self.active_rows())} løp.")

    def reset_music(self):
        self.music_zip = None
        self.music_cache_dir = None
        self.music_file_info = {}
        self.music_files = []
        self.music_durations = {}
        self.music_index = {}
        self.music_status = {}
        self.music_info_var.set("Ingen musikk-zip valgt.")

    def load_music_zip(self, music_zip, match=True):
        # Leses inn i egne variabler og erstatter forrige musikk først når hele zip-en er lest.
        music_zip = Path(music_zip)
        files, index, durations, file_info, status = [], {}, {}, {}, {}
        self.log(f"Leser musikk-zip: {music_zip.name}")
        try:
            with zipfile.ZipFile(music_zip, "r") as mz:
                for entry in mz.infolist():
                    if entry.is_dir():
                        continue
                    if PurePosixPath(entry.filename).name.lower() == "music-status.csv":
                        text = mz.read(entry).decode("utf-8-sig", errors="replace")
                        status = parse_music_status_csv(text)
                        continue
                    declared_suffix = Path(entry.filename).suffix.lower()
                    if declared_suffix not in AUDIO_SUFFIXES:
                        continue
                    data = mz.read(entry)
                    audio_format, detected_suffix = detect_audio_format(data)
                    format_mismatch = bool(detected_suffix and declared_suffix != detected_suffix)
                    files.append(entry.filename)
                    index[entry.filename] = parse_music_path(entry.filename)
                    durations[entry.filename] = get_audio_duration(data, audio_format)
                    file_info[entry.filename] = {
                        "format": audio_format,
                        "suffix": detected_suffix or declared_suffix,
                        "declared_suffix": declared_suffix,
                        "format_mismatch": format_mismatch,
                        "playable": audio_format != "Ukjent",
                    }
                    if audio_format == "Ukjent":
                        self.log(f"ADVARSEL: Ukjent lydformat: {entry.filename}")
                    elif format_mismatch:
                        self.log(
                            f"ADVARSEL: {entry.filename} har filendelsen "
                            f"{declared_suffix}, men er egentlig {audio_format}."
                        )
        except Exception as exc:
            self.log(f"Kunne ikke lese musikk-zip: {music_zip} ({exc})")
            messagebox.showerror("Feil", f"Kunne ikke lese musikk-zip:\n{exc}")
            return False
        self.reset_music()
        self.music_zip = music_zip
        self.music_files, self.music_index, self.music_durations = files, index, durations
        self.music_file_info, self.music_status = file_info, status
        warnings = sum(1 for info in self.music_file_info.values() if info.get("format_mismatch"))
        errors = sum(1 for info in self.music_file_info.values() if not info.get("playable"))
        self.log(
            f"Lydkontroll: {len(self.music_files)} filer, {warnings} med feil filendelse, "
            f"{errors} med ukjent format."
        )
        if self.music_status:
            self.log(f"Leste music-status.csv: {len(self.music_status)} program.")
        elif any(self.music_index.values()):
            self.log("Fant ikke music-status.csv i musikk-zip.")
        self.music_info_var.set(f"{self.music_zip.name} · {len(self.music_files)} lydfiler")
        if match:
            self.match_all_music()
        return True

    def find_music_file(self, row, used_files, fuzzy=True):
        # Returnerer (fil, usikker).
        if not self.music_files:
            return "", False
        if any(self.music_index.values()):
            return find_music_for_row(row, self.music_index, used_files, fuzzy)
        # Eldre musikk-zip uten musikktjenestens filnavn: leter etter navnet i filnavnet.
        return self.match_music_file(row, self.music_files, used_files, fuzzy)

    def apply_music_match(self, row, used_files, fuzzy=True):
        matched, uncertain = self.find_music_file(row, used_files, fuzzy)
        info = self.music_file_info.get(matched, {}) if matched else {}
        seconds = self.music_durations.get(matched) if matched else None
        row["Musikk"] = "ok" if matched else "mangler"
        row["MusikkFil"] = matched
        row["MusikkUsikker"] = uncertain
        row["MusikkStatus"] = (music_status_text(row, matched, self.music_status, uncertain)
                               if self.music_files else "")
        row["MusikkFormat"] = info.get("format", "")
        row["MusikkFormatAvvik"] = info.get("format_mismatch", False)
        row["MusikkSpillbar"] = info.get("playable", bool(matched))
        if matched and seconds is None:
            row["MusikkTid"] = "Klarer ikke å hente tid"
            row["MusikkSek"] = ""
        else:
            row["MusikkTid"] = format_duration(seconds)
            row["MusikkSek"] = int(round(seconds)) if seconds else ""

    def match_rows(self, rows, used):
        # Eksakte navn kobles for alle først, så løsere treff ikke tar filen fra en med eksakt navn.
        for row in rows:
            self.apply_music_match(row, used, fuzzy=False)
        for row in rows:
            if not row.get("MusikkFil"):
                self.apply_music_match(row, used)

    def match_all_music(self):
        used = set()
        self.match_rows(self.rows, used)
        active = self.active_rows()
        missing = [r for r in active if not r.get("MusikkFil")]
        for row in missing:
            self.log(f"Mangler musikk: {run_name(row)} ({row.get('MusikkStatus')})")
        uncertain = [r for r in active if r.get("MusikkUsikker")]
        for row in uncertain:
            self.log(f"ADVARSEL: Usikker musikkobling, sjekk lydfilen: {run_name(row)} → {row.get('MusikkFil')}")
        for filename in self.music_files:
            if filename not in used:
                self.log(f"Lydfil uten utøver: {filename}")
        self.log(f"Musikk: {len(active) - len(missing)} av {len(active)} løp har lydfil"
                 f"{f', {len(uncertain)} usikre' if uncertain else ''}.")

    def match_missing_music(self):
        if not self.music_files:
            return
        used = {r.get("MusikkFil") for r in self.rows if r.get("MusikkFil")}
        self.match_rows([r for r in self.rows if not r.get("MusikkFil")], used)

    def match_music_file(self, row, music_files, used_files, fuzzy=True):
        given = row.get("GivenName")
        family = row.get("FamilyName")
        given_tokens = tokenize_name(given)
        family_tokens = tokenize_name(family)
        if not family_tokens:
            return "", False

        def family_match(fname):
            hay = normalize_text(fname)
            return all(token in hay for token in family_tokens) or family_tokens[0] in hay

        def given_match(fname):
            if not given_tokens:
                return False
            hay = normalize_text(fname)
            return given_tokens[0] in hay

        # Pass 1: require both family + given (if given exists). Flere treff = usikker kobling.
        if given_tokens:
            hits = [f for f in music_files if family_match(f) and given_match(f)]
            free = [f for f in hits if f not in used_files]
            if free:
                used_files.add(free[0])
                return free[0], len(hits) > 1

        # Pass 2 (fuzzy): bare etternavn. Regnes alltid som usikker, f.eks. ved søsken.
        if fuzzy:
            for fname in music_files:
                if fname not in used_files and family_match(fname):
                    used_files.add(fname)
                    return fname, True
        return "", False

    # -------------------- Musikktabell --------------------
    def refresh_table(self):
        selected = self.tree.selection()
        selected_index = self.tree.index(selected[0]) if selected else None
        self.tree.delete(*self.tree.get_children())
        has_music = bool(self.music_files)
        rows_values = []
        nr = 1
        for row in self.rows:
            cancelled = is_cancelled(row.get("Påmelding", ""))
            missing = has_music and not row.get("MusikkFil")
            if cancelled:
                tags = ("cancelled",)
            elif missing:
                tags = ("missing_music",)
            elif has_music and not row.get("MusikkSpillbar", True):
                tags = ("audio_error",)
            elif row.get("MusikkFormatAvvik") or row.get("MusikkUsikker"):
                tags = ("audio_warning",)
            else:
                tags = ()
            if row.get("LopGjennomfort"):
                tags = tags + ("completed",)
            values = (
                "" if cancelled else nr,
                "☑" if row.get("LopGjennomfort") else "☐",
                row.get("StartTid", ""),
                display_name(row),
                row.get("Organisation", ""),
                row.get("Klasse", ""),
                SEGMENT_NAMES.get(row.get("Segment") or "FS", ""),
                row.get("Påmelding", ""),
                (row.get("MusikkStatus") or "Mangler musikk") if missing else row.get("MusikkStatus", ""),
                PurePosixPath(row.get("MusikkFil") or "").name,
                row.get("MusikkFormat", ""),
                row.get("MusikkTid", ""),
            )
            if not cancelled:
                nr += 1
            rows_values.append(values)
            self.tree.insert("", "end", values=values, tags=tags)
        self.autosize_columns(rows_values)
        items = self.tree.get_children()
        if selected_index is not None and selected_index < len(items):
            self.tree.selection_set(items[selected_index])

    def autosize_columns(self, rows_values):
        font = tkfont.nametofont("TkDefaultFont")
        columns = list(self.tree["columns"])
        padding = 24
        max_widths = {}
        for col in columns:
            heading = self.tree.heading(col).get("text", "")
            max_widths[col] = font.measure(str(heading)) + padding
        for row in rows_values:
            for col, value in zip(columns, row):
                width = font.measure(str(value)) + padding
                if width > max_widths.get(col, 0):
                    max_widths[col] = width
        for col in columns:
            self.tree.column(col, width=max_widths.get(col, 80), stretch=col == "musikknavn")

    def update_visible_table_columns(self, changed_column=None):
        visible_columns = tuple(
            column
            for column in self.table_columns
            if self.table_column_vars[column].get()
        )
        if not visible_columns:
            if changed_column:
                self.table_column_vars[changed_column].set(True)
            messagebox.showinfo(
                "Tabellvisning",
                "Minst én kolonne må være synlig.",
            )
            return
        self.tree.configure(displaycolumns=visible_columns)

    def show_all_table_columns(self):
        for variable in self.table_column_vars.values():
            variable.set(True)
        self.tree.configure(displaycolumns=self.table_columns)

    def get_tree_column_name(self, x_position):
        display_column = self.tree.identify_column(x_position)
        if not display_column:
            return None
        display_index = int(display_column[1:]) - 1
        columns = list(self.tree["columns"])
        display_columns = list(self.tree["displaycolumns"])
        if display_columns == ["#all"]:
            display_columns = columns
        if display_index < 0 or display_index >= len(display_columns):
            return None
        return display_columns[display_index]

    def on_tree_checkbox_click(self, event):
        if self.tree.identify("region", event.x, event.y) != "cell":
            return None
        if self.get_tree_column_name(event.x) != "lop_gjennomfort":
            return None

        item_id = self.tree.identify_row(event.y)
        if not item_id:
            return "break"
        row_index = self.tree.index(item_id)
        if row_index < 0 or row_index >= len(self.rows):
            return "break"

        row = self.rows[row_index]
        self.tree.selection_set(item_id)
        self.tree.focus(item_id)
        completed = bool(row.get("LopGjennomfort"))
        name = (
            row.get("NavnFraIsonen")
            or row.get("PrintName")
            or "denne deltakeren"
        )
        if row.get("Program"):
            name = f"{name} ({row['Program'].lower()})"
        if completed:
            question = f"Vil du fjerne markeringen som gjennomført for {name}?"
        else:
            question = f"Vil du markere løpet som gjennomført for {name}?"
        if not messagebox.askyesno(
            "Bekreft løpsstatus",
            question,
            parent=self.root,
        ):
            return "break"

        row["LopGjennomfort"] = not completed
        self.dirty = True
        status = "gjennomført" if not completed else "ikke gjennomført"
        self.log(f"Løpsstatus for {name}: {status}.")
        self.refresh_table()
        items = self.tree.get_children()
        if row_index < len(items):
            self.tree.selection_set(items[row_index])
            self.tree.focus(items[row_index])
        return "break"

    def on_tree_double_click(self, event):
        if self.tree.identify("region", event.x, event.y) != "cell":
            return
        if self.get_tree_column_name(event.x) == "lop_gjennomfort":
            return
        item_id = self.tree.identify_row(event.y)
        if not item_id:
            return
        self.play_mp3_file(self.rows[self.tree.index(item_id)].get("MusikkFil", ""))

    # -------------------- Avspilling --------------------
    def get_cached_mp3_path(self, filename):
        if not filename:
            messagebox.showinfo("Info", "Ingen lydfil registrert på denne raden.")
            return None
        if not self.music_zip:
            messagebox.showerror("Feil", "Fant ingen musikk-zip.")
            return None
        try:
            with zipfile.ZipFile(self.music_zip, "r") as mz:
                try:
                    data = mz.read(filename)
                except KeyError:
                    messagebox.showerror("Feil", f"Fant ikke lydfil i zip: {filename}")
                    return None
        except Exception as exc:
            messagebox.showerror("Feil", f"Kunne ikke lese musikk-zip: {exc}")
            return None

        if not self.music_cache_dir:
            base = Path(tempfile.gettempdir()) / "fms_gui_music_cache"
            base.mkdir(parents=True, exist_ok=True)
            safe_zip = sanitize_filename(self.music_zip.stem)
            self.music_cache_dir = base / safe_zip
            self.music_cache_dir.mkdir(parents=True, exist_ok=True)

        audio_info = self.music_file_info.get(filename)
        if not audio_info:
            audio_format, detected_suffix = detect_audio_format(data)
            audio_info = {
                "format": audio_format,
                "suffix": detected_suffix or Path(filename).suffix.lower(),
            }
            self.music_file_info[filename] = audio_info

        safe_stem = sanitize_filename(Path(filename).stem)
        content_hash = hashlib.sha256(data).hexdigest()[:12]
        safe_name = f"{safe_stem}_{content_hash}{audio_info.get('suffix', '')}"
        out_path = self.music_cache_dir / safe_name
        try:
            if out_path.is_file() and out_path.stat().st_size == len(data):
                return out_path
        except OSError:
            pass
        try:
            out_path.write_bytes(data)
        except Exception as exc:
            messagebox.showerror("Feil", f"Kunne ikke skrive lydfil: {exc}")
            return None
        return out_path

    def ensure_audio_backend(self):
        if self.audio_ready:
            return True
        try:
            import pygame
        except Exception as exc:
            self.log(f"Intern lydmotor (pygame) er ikke tilgjengelig: {exc}")
            return False
        try:
            pygame.mixer.init()
        except Exception as exc:
            self.log(f"Kunne ikke starte intern lydmotor: {exc}")
            return False
        self.audio_backend = pygame
        self.audio_ready = True
        return True

    def release_internal_audio(self):
        if not self.audio_ready:
            return
        try:
            self.audio_backend.mixer.music.stop()
        except Exception:
            pass
        try:
            self.audio_backend.mixer.music.unload()
        except Exception:
            pass

    def set_external_playback_status(self, filename, player_name):
        self.current_track = filename
        self.external_playback = True
        self.is_paused = False
        self.play_pause_text.set("Spill")
        row = self.find_row_by_mp3(filename)
        if row:
            title = row.get("NavnFraIsonen") or row.get("PrintName") or ""
            self.player_track_var.set(f"{title} - {filename}".strip(" -"))
        else:
            self.player_track_var.set(filename)
        self.wave_peaks = []
        self.wave_progress = 0.0
        self.draw_waveform()
        self.player_time_var.set(player_name)
        self.player_total_var.set("")

    def stop_running_playback(self):
        # Stopper VLC og intern avspilling før et nytt spor starter, så to spor aldri overlapper.
        # Windows-standardspilleren (reserve) kan vi ikke styre.
        if self.vlc_process and self.vlc_process.poll() is None:
            try:
                self.vlc_process.terminate()
            except Exception:
                pass
        self.vlc_process = None
        self.external_playback = False
        self.release_internal_audio()

    def start_vlc_playback(self, path, filename, fallback_to_default=True):
        self.stop_running_playback()
        vlc_path = find_vlc_executable()
        if vlc_path:
            try:
                self.vlc_process = subprocess.Popen(
                    [vlc_path, "--play-and-exit", "--no-video", str(path)]
                )
                self.set_external_playback_status(filename, "VLC")
                self.log(f"Spiller med VLC: {filename}")
                return True
            except Exception as exc:
                self.log(f"Kunne ikke starte VLC: {exc}")

        if fallback_to_default:
            try:
                os.startfile(str(path))
                self.set_external_playback_status(filename, "Ekstern avspiller")
                self.log(f"Spiller med Windows-standardspiller: {filename}")
                return True
            except Exception as exc:
                self.log(f"Kunne ikke starte ekstern avspiller: {exc}")

        messagebox.showerror(
            "Feil",
            "Kunne ikke spille av lydfilen. VLC ble ikke funnet eller kunne ikke startes.",
        )
        return False

    def start_playback(self, path, filename):
        self.stop_running_playback()
        if not self.ensure_audio_backend():
            self.log("Intern lydmotor er ikke tilgjengelig. Prøver VLC.")
            return self.start_vlc_playback(path, filename)
        row = self.find_row_by_mp3(filename)
        self.current_duration = safe_int(row.get("MusikkSek")) if row else 0
        try:
            self.release_internal_audio()
            self.audio_backend.mixer.music.load(str(path))
            self.audio_backend.mixer.music.play()
        except Exception as exc:
            self.log(f"Kunne ikke spille av med intern avspiller: {exc}")
            self.release_internal_audio()
            self.log("Prøver VLC som reserveavspiller.")
            return self.start_vlc_playback(path, filename)
        self.current_track = filename
        self.is_paused = False
        self.external_playback = False
        self.play_offset = 0.0
        self.wave_peaks, length = self.load_waveform(path)
        if not self.current_duration:
            self.current_duration = length
        self.wave_progress = 0.0
        self.draw_waveform()
        self.play_pause_text.set("Pause")
        if row:
            title = row.get("NavnFraIsonen") or row.get("PrintName") or ""
            self.player_track_var.set(f"{title} - {filename}".strip(" -"))
        else:
            self.player_track_var.set(filename)
        self.log(f"Spiller: {filename}")
        return True

    def play_mp3_file(self, filename):
        path = self.get_cached_mp3_path(filename)
        if not path:
            return
        if self.use_external_player_var.get():
            self.start_vlc_playback(path, filename)
            return
        self.start_playback(path, filename)

    def get_selected_mp3_filename(self):
        selected = self.tree.selection()
        if not selected:
            return ""
        return self.rows[self.tree.index(selected[0])].get("MusikkFil", "")

    def toggle_play_pause(self):
        if self.use_external_player_var.get():
            messagebox.showinfo(
                "Info",
                "Bruker VLC direkte. Pause/fortsett styres i VLC.",
            )
            return
        if self.external_playback:
            messagebox.showinfo(
                "Info",
                "Spiller med reserveavspiller. Pause/fortsett styres der.",
            )
            return
        if self.audio_ready and self.is_paused:
            self.audio_backend.mixer.music.unpause()
            self.is_paused = False
            self.play_pause_text.set("Pause")
            self.log("Fortsetter avspilling.")
            return
        if self.audio_ready and self.audio_backend.mixer.music.get_busy():
            self.audio_backend.mixer.music.pause()
            self.is_paused = True
            self.play_pause_text.set("Spill")
            self.log("Pause.")
            return
        filename = self.get_selected_mp3_filename()
        if not filename:
            messagebox.showinfo("Info", "Velg en rad med lydfil først.")
            return
        path = self.get_cached_mp3_path(filename)
        if not path:
            return
        self.start_playback(path, filename)

    def stop_playback(self):
        if self.use_external_player_var.get():
            if self.vlc_process and self.vlc_process.poll() is None:
                self.vlc_process.terminate()
                self.vlc_process = None
                self.external_playback = False
                self.play_pause_text.set("Spill")
                self.log("Stoppet VLC-avspilling.")
            else:
                messagebox.showinfo("Info", "Stopp avspillingen i VLC.")
            return
        if self.external_playback:
            if self.vlc_process and self.vlc_process.poll() is None:
                self.vlc_process.terminate()
                self.vlc_process = None
                self.external_playback = False
                self.play_pause_text.set("Spill")
                self.log("Stoppet VLC-avspilling.")
            else:
                messagebox.showinfo("Info", "Stopp avspillingen i den eksterne spilleren.")
            return
        if not self.ensure_audio_backend():
            return
        self.release_internal_audio()
        self.current_track = ""
        self.is_paused = False
        self.external_playback = False
        self.play_offset = 0.0
        self.play_pause_text.set("Spill")
        self.log("Stoppet avspilling.")

    def find_row_by_mp3(self, filename):
        for row in self.rows:
            if row.get("MusikkFil") == filename:
                return row
        return None

    def on_close(self):
        if not self.confirm_discard():
            return
        if self.vlc_process and self.vlc_process.poll() is None:
            try:
                self.vlc_process.terminate()
            except Exception:
                pass
        if self.audio_ready:
            self.release_internal_audio()
            try:
                self.audio_backend.mixer.quit()
            except Exception:
                pass
        self.root.destroy()

    def update_player_ui(self):
        if self.external_playback:
            if self.vlc_process and self.vlc_process.poll() is not None:
                self.vlc_process = None
                self.current_track = ""
                self.external_playback = False
                self.play_pause_text.set("Spill")
                self.player_time_var.set("0:00")
                self.player_total_var.set("0:00")
            self.root.after(500, self.update_player_ui)
            return
        elapsed = 0
        if self.audio_ready and self.current_track:
            elapsed = self.play_offset
            pos_ms = self.audio_backend.mixer.music.get_pos()
            if pos_ms and pos_ms > 0:
                elapsed += pos_ms / 1000.0
            if not self.is_paused and not self.audio_backend.mixer.music.get_busy():
                self.release_internal_audio()
                self.current_track = ""
                self.play_offset = 0.0
                elapsed = 0
                self.play_pause_text.set("Spill")
        duration = self.current_duration or 0
        if duration > 0:
            elapsed = min(elapsed, duration)
            self.player_time_var.set(format_duration(elapsed))
            self.player_total_var.set(format_duration(duration))
            self.set_wave_progress(elapsed / duration)
        elif self.player_enabled:
            self.player_time_var.set("0:00")
            self.player_total_var.set("0:00")
            self.set_wave_progress(0.0)
        self.root.after(200, self.update_player_ui)

    def on_use_external_player_toggle(self):
        self.player_enabled = not self.use_external_player_var.get()
        if not self.player_enabled:
            self.player_time_var.set("VLC direkte")
            self.player_total_var.set("")
        elif not self.external_playback:
            self.player_time_var.set("0:00")
            self.player_total_var.set("0:00")
        for name in self.player_buttons:
            self.draw_player_button(name)

    # -------------------- Avspillerens tegning og spoling --------------------
    def draw_player_button(self, name, hover=False):
        canvas = self.player_buttons[name]
        size = int(canvas["width"])
        c = size / 2
        canvas.delete("all")
        big = name == "play"
        if big:
            fill = ACCENT_HOVER if hover and self.player_enabled else ACCENT
            outline, icon = fill, ACCENT_TEXT
        else:
            fill = BUTTON_HOVER if hover and self.player_enabled else FIELD
            outline, icon = LINE, ACCENT
        if not self.player_enabled:
            fill, outline, icon = PANEL, LINE, TEXT_DISABLED
        canvas.create_oval(1, 1, size - 2, size - 2, fill=fill, outline=outline, width=1.5)
        u = size / 10
        if name == "play" and self.play_pause_text.get() == "Pause":
            for x in (c - 1.6 * u, c + 0.6 * u):
                canvas.create_rectangle(x, c - 1.6 * u, x + u, c + 1.6 * u, fill=icon, outline="")
        elif name == "play":
            canvas.create_polygon(c - 1.1 * u, c - 1.7 * u, c + 1.7 * u, c, c - 1.1 * u, c + 1.7 * u,
                                  fill=icon, outline="")
        elif name == "stop":
            canvas.create_rectangle(c - 1.3 * u, c - 1.3 * u, c + 1.3 * u, c + 1.3 * u, fill=icon,
                                    outline="")
        else:
            # « og »: to vinkler, speilet for spoling bakover.
            d = -1 if name == "back" else 1
            for dx in (-1.1 * u, 0.9 * u):
                x = c + d * dx
                canvas.create_line(x - d * u, c - 1.3 * u, x + d * 0.3 * u, c, x - d * u, c + 1.3 * u,
                                   fill=icon, width=1.8, capstyle="round", joinstyle="round")

    def load_waveform(self, path):
        # Leser lydfilen med pygame og lager toppverdier til lydbølgen. Tar ~0,1 s per fil.
        key = str(path)
        if key not in self.wave_cache:
            peaks, length = [], 0
            try:
                sound = self.audio_backend.mixer.Sound(key)
                length = sound.get_length()
                raw = sound.get_raw()
                samples = memoryview(raw)[: len(raw) // 2 * 2].cast("h")
                bucket = max(1, len(samples) // WAVE_BARS)
                stride = max(1, bucket // 400) | 1
                for i in range(WAVE_BARS):
                    part = samples[i * bucket:(i + 1) * bucket:stride]
                    peaks.append(max(max(part), -min(part)) if len(part) else 0)
                top = max(peaks) or 1
                peaks = [p / top for p in peaks]
            except Exception as exc:
                self.log(f"Kunne ikke lage lydbølge: {exc}")
            self.wave_cache[key] = (peaks, length)
        return self.wave_cache[key]

    def draw_waveform(self):
        canvas = self.wave_canvas
        canvas.delete("all")
        width, height = canvas.winfo_width(), int(canvas["height"])
        mid = height / 2
        if not self.wave_peaks:
            canvas.create_line(0, mid, width, mid, fill=LINE, width=2)
            self.wave_bars = []
            return
        step = width / len(self.wave_peaks)
        bar = max(1.0, step * 0.55)
        self.wave_bars = []
        for i, peak in enumerate(self.wave_peaks):
            h = max(2.0, peak * (height - 4) / 2)
            x = i * step
            self.wave_bars.append(canvas.create_rectangle(x, mid - h, x + bar, mid + h, outline=""))
        self.wave_drawn = -1
        self.set_wave_progress(self.wave_progress)

    def set_wave_progress(self, fraction):
        self.wave_progress = fraction
        bars = getattr(self, "wave_bars", [])
        played = int(round(fraction * len(bars)))
        if played == getattr(self, "wave_drawn", -1):
            return
        for i, item in enumerate(bars):
            self.wave_canvas.itemconfigure(item, fill=ACCENT if i < played else LINE)
        self.wave_drawn = played

    def on_waveform_click(self, event):
        width = self.wave_canvas.winfo_width()
        if width > 0 and self.current_duration:
            self.seek_to(event.x / width * self.current_duration)

    def skip_seconds(self, seconds):
        if self.current_track and not self.external_playback:
            elapsed = self.play_offset + max(0, self.audio_backend.mixer.music.get_pos()) / 1000.0
            self.seek_to(elapsed + seconds)

    def seek_to(self, seconds):
        # Spoler i sporet som spilles. Starter avspillingen på nytt fra posisjonen, og beholder pause.
        if not self.player_enabled or not self.audio_ready or not self.current_track or self.external_playback:
            return
        seconds = max(0.0, min(seconds, max(0.0, (self.current_duration or 0) - 0.5)))
        music = self.audio_backend.mixer.music
        try:
            music.play(start=seconds)
            if self.is_paused:
                music.pause()
        except Exception as exc:
            self.log(f"Kunne ikke spole: {exc}")
            return
        self.play_offset = seconds
        if self.current_duration:
            self.player_time_var.set(format_duration(seconds))
            self.set_wave_progress(seconds / self.current_duration)

    # -------------------- Lagring og eksport --------------------
    def row_key(self, row):
        code = (row.get("ParticipantCode") or "").strip()
        if code:
            return f"code:{code}"
        given = (row.get("GivenName") or "").strip()
        family = (row.get("FamilyName") or "").strip()
        event = (row.get("Event") or "").strip()
        return f"name:{given}|{family}|{event}"

    def load_legacy_order(self, order, path):
        # Gamle rekkefølge-filer: sorterer utøverne som allerede er lastet inn.
        buckets = {}
        for row in self.rows:
            buckets.setdefault(self.row_key(row), []).append(row)
        new_rows = []
        for item in order:
            if isinstance(item, dict):
                if item.get("type") == "pause":
                    continue
                key = item.get("key")
            else:
                key = item
            bucket = buckets.get(key)
            if bucket:
                new_rows.append(bucket.pop(0))
            elif isinstance(item, dict) and (item.get("given") or item.get("family")):
                new_rows.append(new_participant(item.get("given"), item.get("family"), item.get("club")))
        for bucket in buckets.values():
            new_rows.extend(bucket)
        self.rows = new_rows
        self.match_missing_music()
        self.relayout(auto_size=True)
        self.refresh_schedule()
        self.log(f"Lastet gammel rekkefølge-fil: {path}")

    def generate_files(self):
        if not self.rows:
            messagebox.showerror("Feil", "Ingen data lastet.")
            return
        out_dir = self.output_dir()
        base_name = "Deltakerliste"
        results = {}
        if self.var_excel.get():
            results["Excel"] = generate_excel(self.rows, str(out_dir / f"{base_name}.xlsx"), self.log)
        if self.var_html.get():
            results["HTML"] = generate_html(self.rows, str(out_dir / f"{base_name}.html"), base_name, self.log)
        if self.var_pdf.get():
            results["PDF"] = generate_pdf(self.rows, str(out_dir / f"{base_name}.pdf"), base_name, self.log)
        if not results:
            messagebox.showinfo("Info", "Velg minst ett format (PDF, Excel eller HTML) i Rapporter-menyen.")
            return
        self.report_export(results, out_dir, "Rapport ferdig", "Rapporten er lagret i:")

    def report_export(self, results, out_dir, title, text):
        # results: {formatnavn: lyktes}. Viser mappen bare når noe ble skrevet, og sier fra om det som feilet.
        failed = [name for name, ok in results.items() if not ok]
        if len(failed) == len(results):
            messagebox.showerror("Feil", f"Kunne ikke lage {', '.join(failed)}. Se loggen for detaljer.")
            return
        if failed:
            messagebox.showwarning("Delvis fullført", f"Kunne ikke lage {', '.join(failed)}. Se loggen for detaljer.")
        self.show_folder_link(out_dir, title, text)

    def startliste_title(self):
        # Returnerer (tittel, filnavn uten endelse), eller None etter å ha vist feilen.
        settings = self.read_schedule_settings()
        if isinstance(settings, str):
            messagebox.showerror("Feil", settings)
            return None
        if not self.entries:
            messagebox.showwarning("Info", "Ingen utøvere i tidsplanen.")
            return None
        date_obj = parse_date_ddmmyy(self.start_date_var.get())
        if not date_obj:
            messagebox.showerror("Feil", "Ugyldig dato. Bruk DD.MM.ÅÅ.")
            return None
        location = self.location_var.get().strip()
        title = " ".join(
            part for part in (self.event_name_var.get().strip(), location, format_date_long(date_obj)) if part
        )
        return title, sanitize_filename(f"Startliste_{location}_{date_obj:%Y-%m-%d}")

    def export_excel(self):
        result = self.startliste_title()
        if not result:
            return
        title, base_name = result
        path = filedialog.asksaveasfilename(
            title="Eksporter startliste til Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"{base_name}.xlsx",
            initialdir=str(self.output_dir()),
        )
        if not path:
            return
        if generate_startliste_excel(self.entries, path, title, self.log):
            self.show_folder_link(Path(path).parent, "Startliste eksportert", "Excel-fila er lagret i:")
        else:
            messagebox.showerror("Feil", "Kunne ikke skrive Excel-fila. Er den åpen i Excel?")

    def generate_startliste(self):
        result = self.startliste_title()
        if not result:
            return
        title, base_name = result
        out_dir = self.output_dir()
        results = {
            "Excel": generate_startliste_excel(self.entries, str(out_dir / f"{base_name}.xlsx"), title, self.log),
            "PDF": generate_startliste_pdf(self.entries, str(out_dir / f"{base_name}.pdf"), title, self.log),
        }
        self.report_export(results, out_dir, "Startliste ferdig", "Startlisten er lagret i:")

    def generate_playlist_only(self):
        if not self.music_zip:
            messagebox.showerror("Feil", "Velg en musikk-zip først.")
            return
        active = self.active_rows()
        if not active:
            messagebox.showwarning("Info", "Ingen utøvere i listen.")
            return
        out_dir = self.output_dir()
        ok = generate_vlc_playlist(active, out_dir, self.music_zip.stem, self.music_zip, self.log)
        self.report_export({"spilleliste": ok}, out_dir, "Spilleliste ferdig", "Spillelisten er lagret i:")

    def show_folder_link(self, folder_path, title, text):
        win = tk.Toplevel(self.root, bg=BG)
        win.title(title)
        win.resizable(False, False)
        style_window_frame(win)
        frame = ttk.Frame(win, padding=16)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text=text).pack(anchor="w")
        path_text = str(folder_path)
        link = ttk.Label(frame, text=path_text, foreground=ACCENT, cursor="hand2")
        link.pack(anchor="w", pady=(4, 12))

        def open_folder(event=None):
            try:
                os.startfile(path_text)
            except Exception as exc:
                messagebox.showerror("Feil", f"Kunne ikke åpne mappe: {exc}")

        link.bind("<Button-1>", open_folder)
        btns = ttk.Frame(frame)
        btns.pack(fill="x")
        ttk.Button(btns, text="Åpne mappe", style="Accent.TButton", command=open_folder).pack(side="right")
        ttk.Button(btns, text="Lukk", command=win.destroy).pack(side="right", padx=(0, 8))

    def show_about(self):
        version = get_version()
        now = datetime.now()
        month_year = format_date_long(now.date()).split(" ", 1)[1]

        about = tk.Toplevel(self.root, bg=BG)
        about.title("Om")
        about.resizable(False, False)
        style_window_frame(about)
        frame = ttk.Frame(about, padding=20)
        frame.pack(fill="both", expand=True)

        logo_path = Path(__file__).resolve().parent / "Lil Logo.jpg"
        if logo_path.exists():
            try:
                from PIL import Image, ImageTk

                img = Image.open(logo_path)
                img.thumbnail((180, 180))
                photo = ImageTk.PhotoImage(img)
                logo_label = ttk.Label(frame, image=photo)
                logo_label.image = photo
                logo_label.pack(pady=(0, 8))
            except Exception:
                pass

        ttk.Label(frame, text=APP_NAME, style="Title.TLabel").pack()
        ttk.Label(frame, text="Loddefjord IL Kunstløp", style="Dim.TLabel").pack()
        ttk.Label(
            frame,
            text=(
                "Setter opp tidsplan for stevnet og spiller av musikken underveis. "
                "Leser deltakerlister fra isonen og musikktjenesten, eller lister du lager selv."
            ),
            wraplength=360,
            justify="center",
        ).pack(pady=(10, 2))
        ttk.Label(frame, text=f"Revisjon: {version} • {month_year}", style="Dim.TLabel").pack(pady=(6, 0))


def main():
    root = tk.Tk()
    App(root)
    style_window_frame(root)
    root.update_idletasks()
    width = int(root.winfo_screenwidth() * 0.85)
    height = int(root.winfo_screenheight() * 0.85)
    root.geometry(f"{width}x{height}")
    root.mainloop()


if __name__ == "__main__":
    main()
